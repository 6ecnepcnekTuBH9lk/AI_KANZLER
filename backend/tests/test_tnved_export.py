"""Golden user format and extraction regressions; use the unchanged strict engine."""

import base64
from copy import deepcopy
from types import SimpleNamespace

import pytest
from app.modules.tnved.comments import user_comment
from app.modules.tnved.features import extract_features
from app.modules.tnved.service import run
from openpyxl import load_workbook
from test_tnved_rules import BASELINE, EXPECTED, engine, garment

COMMENTS = {
    "CONFIRMED": "",
    "SHOE_MATERIALS": "Недостаточно данных для кода обуви: укажите материал верха и материал подошвы.",
    "DENIM_AND_USE": "Недостаточно данных для однозначного кода: необходимо подтвердить, изготовлены ли брюки из денима (джинсовой ткани).",
    "PROFESSIONAL_USE": "Недостаточно данных для полного кода: уточните, является ли изделие производственной или профессиональной одеждой.",
    "SHOE_CONSTRUCTION": "Недостаточно данных для полного кода: не указана длина стельки мужских лоферов с верхом из натуральной замши и подошвой из резины.",
    "UNKNOWN_FIBER": "Состав указан недостаточно подробно: «микрофибра» не раскрывает материал волокна (например, полиэстер или полиамид).",
    "KNIT_DENSITY_AND_COLLAR": "Недостаточно данных для полного кода хлопкового джемпера: необходимо уточнить конструкцию и вид вязки изделия.",
}


@pytest.mark.parametrize("name", ["TENCEL", "Tencel", "тенсел"])
@pytest.mark.parametrize("kind,code", [("Пиджак", "6103390000"), ("Брюки", "6103490001")])
def test_tencel_normalization_and_full_branch(name, kind, code):
    features = extract_features(garment(kind, f"68% {name}, 29% вискоза, 3% спандекс"))
    assert features.material == "искусственные"
    assert features.composition == [("искусственные", 97.0), ("синтетические", 3.0)]
    assert not features.unknown_fibers
    assert engine().classify(features)["code"] == code


def test_tencel_does_not_resolve_independent_denim_condition():
    raw = dict(BASELINE["rows"][11])  # Excel row 13.
    raw.update({"Промышленная или профессиональная одежда": "Нет", "Вельвет-корд с разрезным ворсом": "Нет"})
    features = extract_features(raw)
    assert features.material == "хлопок" and not features.unknown_fibers
    result = engine().classify(features)
    assert result["code"] is None and "деним" in result["comment"]
    assert user_comment(features, result) == COMMENTS["DENIM_AND_USE"]


@pytest.mark.parametrize("row", [24, 25, 26, 27])
def test_professional_and_ordinary_branches_stay_unresolved(row):
    result = engine().classify(extract_features(BASELINE["rows"][row - 2]))
    assert result["code"] is None
    pair = {"6203331000", "6203339000"} if row in (24, 26) else {"6203431100", "6203431900"}
    assert pair <= {c["code"] for c in result["candidates"] if c["verdict"] == "unresolved"}


def test_shoe_upper_from_structured_parts():
    features = extract_features(BASELINE["rows"][26])
    assert features.details["материал верха"] == "натуральная замша"
    assert features.details["материал подошвы"] == "резина"
    result = engine().classify(features)
    assert result["code"] is None
    assert not any("материал верха" in m or "материал подошвы" in m for m in result["missing_input"])
    assert "подкладка" in features.parts and "стелька" in features.parts


@pytest.mark.parametrize(
    "composition",
    [
        "натуральная замша",
        "натуральная замша; резина; подошва: резина",
        "натуральная замша или текстиль; подкладка: кожа; подошва: резина",
        "натуральная замша; подкладка: кожа; подкладка: текстиль; подошва: резина",
        "натуральная замша; отделка: кожа; подошва: резина",
        "натуральная замша; стелька: кожа",
        "натуральная замша; подкладка: кожа; подошва: резина; текстиль",
    ],
)
def test_ambiguous_shoe_composition_never_infers_upper(composition):
    features = extract_features({"Номенклатура": "Лоферы", "Пол": "Мужской", "Состав": composition})
    assert "материал верха" not in features.details


def test_user_comment_ignores_excluded_branches_and_preserves_evidence():
    features = extract_features(garment("Джинсы", Трикотаж="Нет"))
    result = {
        "code": None,
        "missing_input": ["исключение альтернативы: деним"],
        "candidates": [{"verdict": "excluded", "missing_input": ["деним"], "evidence": {"sha256": "secret"}}],
    }
    before = deepcopy(result)
    assert "деним" not in user_comment(features, result)
    assert result == before


@pytest.mark.parametrize("styled_trailing", [False, True])
def test_real_57_excel_golden_source_header_and_comments(tmp_path, monkeypatch, styled_trailing):
    source = tmp_path / "regression.xlsx"
    source.write_bytes(base64.b64decode(BASELINE["xlsx_base64"]))
    if styled_trailing:
        workbook = load_workbook(source)
        workbook.active["M1"].number_format = "@"  # Original bug: an existing blank 13th column.
        workbook.active["M58"].number_format = "@"
        workbook.save(source)
        workbook.close()
    monkeypatch.setattr("app.modules.tnved.service.ClassificationEngine", lambda *a: engine())
    _, sections, output, *_ = run(
        [SimpleNamespace(path=source, original_name=source.name)],
        {},
        tmp_path / "out",
        lambda *a: None,
        None,
        client_factory=lambda: SimpleNamespace(close=lambda: None),
    )
    workbook = load_workbook(output)
    sheet = workbook.active
    assert sheet.max_column == 13
    assert sheet["F1"].value == "Код ТНВЭД"
    assert sheet["L1"].value == "Комментарий"
    assert sheet["M1"].value == "Источник"
    for item, target in zip(sections["items"], EXPECTED["rows"], strict=True):
        row = item["row"]
        comment = COMMENTS[target["category"]]
        assert item["code"] == target["code"]
        assert item["user_comment"] == comment
        assert (sheet.cell(row, 12).value or "") == comment
        assert len(comment) <= 180 and comment.count(".") <= 2
        assert (sheet.cell(row, 6).value or None) == target["code"]
        url = f"https://www.alta.ru/tnved/code/{target['code']}/" if target["code"] else ""
        assert item["source_url"] == url and (sheet.cell(row, 13).value or "") == url
        if target["code"]:
            assert sheet.cell(row, 6).data_type == "s"
        else:
            assert item["comment"] != comment and item["candidates"] and item["missing_input"]
    workbook.close()


@pytest.mark.parametrize("formula_source", [False, True])
def test_reexport_clears_stale_source_and_comment_and_preserves_formulas(
    tmp_path, monkeypatch, formula_source
):
    source = tmp_path / "stale.xlsx"
    source.write_bytes(base64.b64decode(BASELINE["xlsx_base64"]))
    workbook = load_workbook(source)
    sheet = workbook.active
    sheet["L1"] = "Комментарий"
    sheet["M1"] = "Источник Alta.ru"
    sheet["L10"] = "Старый технический комментарий с кандидатами"
    sheet["M13"] = '= "Сохранить"' if formula_source else "https://www.alta.ru/ett/gruppa62.html"
    sheet["Z10"] = "Внешние данные"
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr("app.modules.tnved.service.ClassificationEngine", lambda *a: engine())
    _, _, output, *_ = run(
        [SimpleNamespace(path=source, original_name=source.name)],
        {},
        tmp_path / "out",
        lambda *a: None,
        None,
        client_factory=lambda: SimpleNamespace(close=lambda: None),
    )
    workbook = load_workbook(output)
    sheet = workbook.active
    assert not sheet["L10"].value
    source_col = 27 if formula_source else 13
    assert sheet.cell(1, source_col).value == "Источник"
    assert not sheet.cell(13, source_col).value
    assert sheet.cell(10, source_col).value == "https://www.alta.ru/tnved/code/6109902000/"
    assert sheet["Z10"].value == "Внешние данные"
    if formula_source:
        assert sheet["M13"].value == '= "Сохранить"'
    workbook.close()
