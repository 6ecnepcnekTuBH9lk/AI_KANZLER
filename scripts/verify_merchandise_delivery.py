"""Validate downloaded XLSX/persistence and a legacy snapshot in the isolated acceptance DB."""

import hashlib
import json
import shutil
import sqlite3
import sys
from pathlib import Path

import httpx
from merchandise_fixture import ROOT
from openpyxl import load_workbook

sys.path.insert(0, str(ROOT / "backend/tests"))
sys.path.insert(0, str(ROOT))
from app.core.config import Settings
from app.core.database import make_database
from app.shared.models import Job, RunRow
from sqlalchemy import select
from test_merchandise_golden import test_all_exported_values_and_compact_layout

OUT = ROOT / "test-results/merchandise-followup"
REAL = ROOT / "test-results/real-merchandise"


def digest(value):
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()
    ).hexdigest()


def run():
    manifest = json.loads((REAL / "http-verification.json").read_text(encoding="utf-8"))
    payload = json.loads((REAL / "result.json").read_text(encoding="utf-8"))
    path = Path(manifest["output"])
    test_all_exported_values_and_compact_layout((None, (payload["result"], payload["sections"], path)))
    data = OUT / "server"
    database = data / "database/platform.db"
    with sqlite3.connect(f"{database.as_uri()}?mode=ro", uri=True) as db:
        result = json.loads(
            db.execute("SELECT result FROM jobs WHERE id=?", (manifest["job_id"],)).fetchone()[0]
        )
        assert result == payload["result"]
        counts = {}
        for section, rows in payload["sections"].items():
            stored = [
                json.loads(r[0])
                for r in db.execute(
                    "SELECT data FROM analysis_rows WHERE job_id=? AND section=? ORDER BY id",
                    (manifest["job_id"], section),
                )
            ]
            assert stored == rows
            counts[section] = len(stored)
    # The legacy fixture is inserted only into this task's isolated acceptance store.
    assert data.resolve().is_relative_to((ROOT / "test-results").resolve())
    old = json.loads((OUT / "before.json").read_text(encoding="utf-8"))
    legacy_id = "merchandise-followup-legacy"
    legacy_path = OUT / "legacy.xlsx"
    if not legacy_path.exists():
        shutil.copyfile(Path.home() / "Downloads/Анализ продаж приложение.xlsx", legacy_path)
    engine, sessions = make_database(Settings(data, f"sqlite:///{database.as_posix()}"))
    with sessions() as session:
        if not session.get(Job, legacy_id):
            session.add(
                Job(
                    id=legacy_id,
                    module="merchandise",
                    title="Проверка совместимости — прежний результат",
                    status="completed",
                    progress=100,
                    message="Сохранённый результат до доработки",
                    file_ids=[],
                    result=old["result"],
                    output_path=str(legacy_path),
                    output_name=legacy_path.name,
                )
            )
            session.flush()
            session.add_all(
                RunRow(
                    job_id=legacy_id,
                    section=section,
                    item_key=row.get("article") or row.get("category") or str(i),
                    data=row,
                )
                for section, rows in old["sections"].items()
                for i, row in enumerate(rows)
            )
            session.commit()
        before = digest(session.get(Job, legacy_id).result)
    with httpx.Client(base_url="http://127.0.0.1:8001", timeout=60) as client:
        response = client.get(f"/api/merchandise/runs/{legacy_id}")
        response.raise_for_status()
        assert response.json()["result"] == old["result"]
        assert "presentation_version" not in response.json()["result"]
        response = client.get(f"/api/merchandise/runs/{legacy_id}/export")
        assert hashlib.sha256(response.content).digest() == hashlib.sha256(legacy_path.read_bytes()).digest()
    with sessions() as session:
        assert digest(session.get(Job, legacy_id).result) == before
        rows = session.scalars(
            select(RunRow).where(RunRow.job_id == legacy_id, RunRow.section == "articles").order_by(RunRow.id)
        ).all()
        assert [r.data for r in rows] == old["sections"]["articles"]
    engine.dispose()
    wb = load_workbook(path, data_only=False)
    report = {
        "job_id": manifest["job_id"],
        "sqlite_api_exact": True,
        "sections": counts,
        "all_exported_cells_match": True,
        "checked_cells": sum(ws.max_row * ws.max_column for ws in wb),
        "sheets": wb.sheetnames,
        "articles_columns": wb["Артикулы"].max_column,
        "grouped_qa_rows": wb["Качество данных"].max_row - 1,
        "legacy_job_id": legacy_id,
        "legacy_json_unchanged_sha256": before,
        "legacy_export_unchanged": True,
        "input_sha256": manifest["inputs_sha256"],
    }
    wb.close()
    (OUT / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    run()
