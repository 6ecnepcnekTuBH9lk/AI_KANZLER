import math
from datetime import date, timedelta

import pytest
from app.modules.merchandise import decisions, importers, metrics, planning
from app.modules.merchandise.reporting import SHEETS
from app.modules.merchandise.service import run
from app.shared.excel import fraction
from openpyxl import load_workbook

START = date(2026, 9, 1)
END = date(2027, 2, 28)


@pytest.mark.parametrize(
    "text,expected",
    [("ОЗ26_W33", "ОЗ26"), ("Осень-Зима 2026", "ОЗ26"), ("ВЛ27", "ВЛ27"), ("Прогноз 31.08", None)],
)
def test_season_detection(text, expected):
    assert importers.parse_season(text) == expected


def test_rebase_ignores_first_batch(article_plan):
    article_plan["1-я партия"] = 1
    a, _ = planning.build_plan(article_plan, None, 1000, START, END)
    article_plan["1-я партия"] = 10**12
    b, _ = planning.build_plan(article_plan, None, 1000, START, END)
    assert a == b
    assert a["season_units"] == 700
    assert math.isclose(sum(a["daily"].values()), 0.7)


def test_oz26_boundaries_and_daily_shape(article_plan):
    p, _ = planning.build_plan(article_plan, None, 1000, START, END)
    assert min(p["daily"]) == START and max(p["daily"]) == END
    assert len(p["weekly"]) == 26
    assert p["weekly"][0]["week_end"] == "2026-09-06"
    assert p["weekly"][-1]["week_end"] == "2027-02-28"
    pct, units = planning.plan_at(p, 1000, date(2026, 9, 6), START, END)
    assert pct == pytest.approx(0.7 * 0.1 / 30 * 6)
    assert units == pytest.approx(14)
    assert planning.plan_at(p, 1000, date(2026, 8, 31), START, END) == (None, None)
    assert planning.plan_at(p, 1000, date(2027, 3, 10), START, END)[1] == pytest.approx(700)


def test_article_priority_category_fallback_and_missing_plan(article_plan):
    category = {**article_plan, "months": {m: 1 for m in (9, 10, 11, 12, 1, 2)}}
    a, _ = planning.build_plan(article_plan, category, 1000, START, END)
    b, _ = planning.build_plan(article_plan, None, 1000, START, END)
    assert a == b
    assert planning.build_plan(None, category, 1000, START, END)[0]["season_units"] == 700
    article_plan["months"] = {}
    c, qa = planning.build_plan(article_plan, category, 1000, START, END)
    assert c and qa


def test_zero_and_invalid_base(article_plan):
    for base in (0, None, -1):
        assert planning.build_plan(article_plan, None, base, START, END)[0] is None


def test_preseason_boundary_and_postseason(merchandise_fact):
    merchandise_fact["weekly"] = {date(2026, 8, 24): 100, date(2026, 8, 31): 70}
    fact, pre, qa = metrics.sales_window(merchandise_fact, START, date(2026, 9, 6), END)
    assert fact == 60 and pre == 110 and qa
    merchandise_fact["daily"] = {START: 2, START + timedelta(days=1): 3, START - timedelta(days=1): 99}
    fact, pre, qa = metrics.sales_window(merchandise_fact, START, START + timedelta(days=1), END)
    assert (fact, pre) == (5, 99)


def test_partial_week_pace_and_zero(merchandise_fact):
    p, _qa = metrics.pace_metrics(merchandise_fact, date(2026, 10, 1), START, END)
    assert p["partial_week"] and p["wow"] is None
    assert p["avg2"] == 55 and p["avg4"] == 50
    merchandise_fact["weekly"] = {
        date(2026, 9, 7): 0,
        date(2026, 9, 14): 0,
        date(2026, 9, 21): 0,
        date(2026, 9, 28): 0,
    }
    p, _ = metrics.pace_metrics(merchandise_fact, date(2026, 10, 4), START, END)
    assert p["pace"] == 0


def test_required_forecast_cover_st(merchandise_fact, article_plan):
    article_plan["deadline"] = END
    merchandise_fact["weekly"][date(2026, 8, 31)] = 70
    p, _ = planning.build_plan(article_plan, None, 1000, START, END)
    pct, units = planning.plan_at(p, 1000, date(2026, 10, 4), START, END)
    m, _ = metrics.calculate(merchandise_fact, article_plan, p, date(2026, 10, 4), START, END, pct, units)
    assert m["season_fact"] == 280
    assert m["st"] == 0.28
    assert m["execution"] == pytest.approx(280 / units)
    assert m["gap_units"] == pytest.approx(280 - units)
    assert m["gap_pp"] == pytest.approx((0.28 - pct) * 100)
    assert m["pace"] == 55
    assert m["cover"] == pytest.approx(600 / 55)
    assert m["required"] == pytest.approx((700 - 280) / m["weeks_remaining"])
    assert m["forecast"] == pytest.approx(280 + 55 * m["weeks_remaining"])


@pytest.mark.parametrize(
    "execution,expected",
    [
        (1.15, "ХИТ"),
        (1.1499, "В ПЛАНЕ"),
        (0.85, "В ПЛАНЕ"),
        (0.8499, "РИСК"),
        (0.70, "РИСК"),
        (0.6999, "АУТСАЙДЕР"),
        (0, "АУТСАЙДЕР"),
    ],
)
def test_status_boundaries(execution, expected, article_plan):
    m = {
        "season_observation_days": 34,
        "age": 35,
        "norms": {"observation_days": {"value": 28}},
        "execution": execution,
        "deadline": END.isoformat(),
        "pace": 10,
        "forecast": 600 if execution < 0.70 else 800 if execution >= 1.15 else 700,
        "season_plan": 700,
        "season_pace_weeks": 4,
    }
    ops = {k: True for k, l in decisions.OPERATION_ORDER}
    assert decisions.status(m, ops, None, date(2026, 10, 4), START, article_plan) == expected


def test_status_observation_forecast_sizes_preseason(article_plan):
    m = {
        "season_observation_days": 34,
        "age": 35,
        "norms": {"observation_days": {"value": 28}},
        "execution": 1.0,
        "deadline": END.isoformat(),
        "pace": 10,
        "forecast": 600,
        "season_plan": 700,
    }
    ops = {k: True for k, l in decisions.OPERATION_ORDER}
    assert decisions.status(m, ops, None, date(2026, 10, 4), START, article_plan) == "РИСК"
    m["season_observation_days"] = 3
    assert decisions.status(m, ops, None, date(2026, 10, 4), START, article_plan) == "НЕДОСТАТОЧНО ДАННЫХ"
    m["season_observation_days"] = 34
    assert decisions.status(m, ops, None, date(2026, 8, 31), START, article_plan) == "НЕДОСТАТОЧНО ДАННЫХ"
    m["execution"] = 0.5
    ops["sizes_ok"] = False
    assert decisions.status(m, ops, None, date(2026, 10, 4), START, article_plan) == "АУТСАЙДЕР"


def test_second_wave_base_and_post_arrival(merchandise_fact, article_plan):
    merchandise_fact["second_actual"] = date(2026, 12, 1)
    article_plan.update(second_date=date(2026, 12, 1), second_qty=500)
    m = {"base": 1000, "pace": 20, "season_fact": 200, "preseason": 100, "cover": 30}
    wave = decisions.second_wave(merchandise_fact, m, article_plan, date(2026, 10, 1), END)
    assert wave["target70"] == 700 and wave["first_sales"] == 300 and wave["remaining70"] == 400
    assert wave["risk"] and wave["decision"] == "удержать у поставщика"
    wave = decisions.second_wave(merchandise_fact, m, article_plan, date(2026, 12, 2), END)
    assert wave["first_sales"] is None and wave["forecast"] is None


def test_pricing_operations_order_and_no_automatic_discount(merchandise_fact, article_plan):
    m = {"forecast": 200, "season_plan": 700, "cover": 30, "weeks_remaining": 10}
    ops = {k: None for k, l in decisions.OPERATION_ORDER}
    p = decisions.pricing(merchandise_fact, m, article_plan, "РИСК", ops, None, date(2026, 10, 4))
    assert p["decision"] == "наблюдать" and "представленность" in p["reason"]
    ops = {k: True for k, l in decisions.OPERATION_ORDER}
    p = decisions.pricing(merchandise_fact, m, article_plan, "РИСК", ops, None, date(2026, 10, 4))
    assert p["decision"] == "проверить данные" and "предел" in p["reason"]
    m["norms"] = {"max_discount": {"value": 0.4}}
    p = decisions.pricing(merchandise_fact, m, article_plan, "РИСК", ops, None, date(2026, 10, 4))
    assert p["decision"] == "снизить цену" and "40%" in p["change"]
    merchandise_fact["repricing_date"] = date(2026, 9, 15)
    p = decisions.pricing(merchandise_fact, m, article_plan, "РИСК", ops, None, date(2026, 10, 4))
    assert p["decision"] == "наблюдать" and "не подтверждён" in p["reason"]


def test_synthetic_full_export_and_rebase_regression(commercial_files, tmp_path):
    first = run(commercial_files, {}, tmp_path / "first", lambda p, m: None)
    result, sections, path, _, _ = first
    assert result["season"] == "ОЗ26" and result["control_date"] == "2026-10-04"
    a = sections["articles"][0]
    assert a["base"] == 1000 and a["season_plan"] == 700
    assert a["season_fact"] == pytest.approx(190)
    assert a["execution"] == pytest.approx(a["season_fact"] / a["plan_units"])
    wb = load_workbook(path)
    assert wb.sheetnames == list(SHEETS.values())
    headers = [c.value for c in wb["Артикулы"][1]]
    assert "Выполнение плана на текущую дату, %" in headers
    assert not any("1-я партия" in h or "План неделя" in h for h in headers)
    assert all(ws.freeze_panes and ws.auto_filter.ref for ws in wb)
    assert wb["План по неделям"].cell(1, 5).value == "План неделя 36 2026, % накоп."
    wb.close()
    source = load_workbook(commercial_files[0].path)
    source.active.cell(2, 13).value = 1
    source.save(commercial_files[0].path)
    second = run(commercial_files, {}, tmp_path / "second", lambda p, m: None)
    assert first[0] == second[0] and first[1] == second[1]


def test_one_percent_1c_and_percentage_cell():
    assert fraction(1, whole=True) == 0.01
    assert fraction(1, "0%", whole=True) == 1
    assert fraction(0, whole=True) == 0


def test_daily_sales_support_completed_week_pace(merchandise_fact):
    merchandise_fact["weekly"] = {}
    merchandise_fact["daily"] = {d: 2 for d in planning.days(date(2026, 9, 7), date(2026, 10, 4))}
    m, _ = metrics.pace_metrics(merchandise_fact, date(2026, 10, 4), START, END)
    assert m["pace"] == 14 and m["season_pace_weeks"] == 4
    del merchandise_fact["daily"][date(2026, 10, 3)]
    m, _ = metrics.pace_metrics(merchandise_fact, date(2026, 10, 4), START, END)
    assert m["pace"] is None


def test_duplicate_and_negative_fact_are_not_calculated(commercial_files, tmp_path):
    source = load_workbook(commercial_files[2].path)
    source.active.cell(2, 6).value = -2
    source.save(commercial_files[2].path)
    _, sections, *_ = run(commercial_files, {}, tmp_path / "negative", lambda p, m: None)
    assert sections["articles"][0]["status"] == "НЕДОСТАТОЧНО ДАННЫХ"
    assert sections["articles"][0]["forecast"] is None
    source.active.append([cell.value for cell in source.active[2]])
    source.save(commercial_files[2].path)
    _, sections, *_ = run(commercial_files, {}, tmp_path / "duplicate", lambda p, m: None)
    assert len(sections["articles"]) == 1 and sections["articles"][0]["base"] is None


def test_early_plan_deadline_is_not_extended(article_plan):
    article_plan["deadline"] = date(2026, 12, 31)
    trajectory, warnings = planning.build_plan(article_plan, None, 1000, START, END)
    assert trajectory and any("после нормативного" in w for w in warnings)
    assert metrics.deadline_for(article_plan, date(2026, 9, 1), END) == date(2026, 12, 31)


def test_high_preseason_pace_does_not_confirm_hit(article_plan):
    m = {
        "season_observation_days": 6,
        "age": 35,
        "norms": {"observation_days": {"value": 28}},
        "execution": 1.2,
        "deadline": END.isoformat(),
        "pace": 100,
        "forecast": 3000,
        "season_plan": 700,
        "season_pace_weeks": 0,
    }
    operations = {key: True for key, _ in decisions.OPERATION_ORDER}
    assert (
        decisions.status(m, operations, None, date(2026, 9, 6), START, article_plan) == "НЕДОСТАТОЧНО ДАННЫХ"
    )
    m["season_observation_days"] = 34
    assert decisions.status(m, operations, None, date(2026, 10, 4), START, article_plan) == "В ПЛАНЕ"
