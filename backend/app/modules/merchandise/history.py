from datetime import date, datetime
from functools import lru_cache

from app.core.config import RESOURCES
from app.shared.excel import norm, number, read_tables
from openpyxl import load_workbook

ROOT = RESOURCES / "kanzler-merchandise-performance-controller"


@lru_cache(maxsize=1)
def historical_index():
    result = {"categories": {}, "articles": {}, "final": {}}
    for name, key in [
        ("ОЗ25 продажи по неделям (по группам-видам).xlsx", "categories"),
        ("ОЗ25 продажи по неделям (по-артикульно).xlsx", "articles"),
    ]:
        wb = load_workbook(ROOT / name, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows)
        date_cols = [
            (i, v.date() if isinstance(v, datetime) else v)
            for i, v in enumerate(headers)
            if isinstance(v, (date, datetime))
        ]
        next(rows)
        for row in rows:
            identity = (norm(row[1]), norm(row[2])) if key == "categories" else str(row[3] or "")
            weekly = {d.isoformat(): number(row[i]) or 0 for i, d in date_cols}
            total = sum(weekly.values())
            result[key][identity] = {
                "source": name,
                "total": total,
                "weekly_shares": {d: q / total if total else None for d, q in weekly.items()},
            }
        wb.close()
    tables = read_tables(ROOT / "Анализ продаж ОЗ 25 на  02.03.2026.xlsx")
    for table in tables:
        if table.col("Артикул") is None:
            continue
        for _, r in table.rows:
            a = str(table.get(r, "Артикул") or "")
            if not a.startswith("5A"):
                continue
            key = (norm(table.get(r, "Вид номенклатуры")), norm(table.get(r, "Вид ассортимента")))
            base = number(table.get(r, "Начальный остаток"))
            qty = number(table.get(r, "Продано за 12 мес"))
            bucket = result["final"].setdefault(key, {"base": 0, "sales": 0, "articles": 0})
            if base is not None and qty is not None:
                bucket["base"] += base
                bucket["sales"] += qty
                bucket["articles"] += 1
    return result


def context(fact, control):
    index = historical_index()
    key = (norm(fact["kind"]), norm(fact["assortment"]))
    cat = index["categories"].get(key)
    if not cat:
        return {
            "text": "Не найдено точное историческое соответствие категории и вида ассортимента.",
            "weekly_shares": {},
            "comparable_article": None,
        }
    final = index["final"].get(key)
    details = f"ОЗ25: относительная недельная кривая категории из «{cat['source']}»."
    if final and final["base"]:
        details += f" Итоговый срез: {final['articles']} артикулов, продажи к начальной базе {final['sales'] / final['base']:.1%}; период среза шире официального сезонного окна."
    details += (
        " История — контекст; различия в глубине, ширине матрицы, датах входа и цене ограничивают сравнение."
    )
    comparable = index["articles"].get(fact["article"])
    return {
        "text": details,
        "weekly_shares": cat["weekly_shares"],
        "comparable_article": comparable,
        "final_snapshot": final,
        "source_article_count": len(index["articles"]),
    }
