"""One-time immutable input/before capture; never use snapshots in production."""

import base64
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "backend/tests/fixtures/tnved"
DOWNLOADS = Path.home() / "Downloads"
NAME = "Артикулы_коллеции_ВЛ_27_Нет_кодов_ТНВЭД_2"


def main():
    DEST.mkdir(parents=True, exist_ok=True)
    target = DEST / "baseline.json"
    if target.exists():
        raise SystemExit("Baseline already exists; refusing to rewrite.")
    source = DOWNLOADS / (NAME + ".xlsx")
    blob = source.read_bytes()
    book = load_workbook(source, data_only=False)
    rows = list(book.active.values)
    originals = [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]
    previous = []
    for name in (NAME + " (1).xlsx", NAME + "_с_кодами_ТНВЭД.xlsx"):
        path = DOWNLOADS / name
        wb = load_workbook(path, data_only=False)
        values = list(wb.active.values)
        previous.append(
            {
                "name": name,
                "provenance": "UNCONFIRMED",
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "rows": [dict(zip(values[0], row, strict=True)) for row in values[1:]],
            }
        )
        wb.close()
    before = json.loads((ROOT / "test-results/real-tnved/result.json").read_text(encoding="utf-8"))
    payload = {
        "commit": "787c689",
        "name": source.name,
        "sha256": hashlib.sha256(blob).hexdigest(),
        "xlsx_base64": base64.b64encode(blob).decode(),
        "rows": originals,
        "app_before": before,
        "previous_files": previous,
    }
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Captured", len(originals), "rows", payload["sha256"])
    for i, row in enumerate(originals, 2):
        print(i, json.dumps(row, ensure_ascii=False))
    book.close()


if __name__ == "__main__":
    main()
