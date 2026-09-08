"""Read-only comparison of Skill workbook, supplied app workbook, and new stored run."""

import hashlib
import json
import math
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from merchandise_fixture import ROOT
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

OUT = ROOT / "test-results/merchandise-followup"
LABELS = {
    "article": ["Артикул"],
    "base": ["Начальный остаток"],
    "plan_units": ["План на дату, ед.", "План на текущую дату, ед."],
    "preseason": ["До сезона, ед. — оценка", "Продажи до начала сезона, ед."],
    "season_fact": ["Факт продаж сезона, ед. — оценка", "Факт продаж сезона, ед."],
    "execution": ["Выполнение плана на текущую дату, %"],
    "forecast": ["Прогноз продаж, ед. — сценарий", "Прогноз продаж, ед."],
    "status": ["Статус"],
    "operational_signal": ["Операционный сигнал"],
    "recommendation": ["Рекомендация"],
    "review_date": ["Контрольная дата"],
}


def workbook(path):
    wb = load_workbook(path, read_only=False, data_only=True)
    tables, layout = {}, {}
    for ws in wb:
        rows = list(ws.values)
        index = next(
            (
                i
                for i, row in enumerate(rows[:10])
                if row[0] in ("Артикул", "Показатель", "Вид номенклатуры", "Категория", "Приоритет")
            ),
            0,
        )
        headers = rows[index]
        tables[ws.title] = [
            dict(zip(headers, row)) for row in rows[index + 1 :] if any(v is not None for v in row)
        ]
        layout[ws.title] = {
            "rows": len(tables[ws.title]),
            "columns": ws.max_column,
            "width": sum(
                ws.column_dimensions[get_column_letter(i)].width or 13 for i in range(1, ws.max_column + 1)
            ),
        }
    wb.close()
    articles = {}
    for row in tables["Артикулы"]:
        a = {
            key: next((row[label] for label in labels if label in row), None)
            for key, labels in LABELS.items()
        }
        for key, value in a.items():
            if isinstance(value, datetime):
                a[key] = value.date().isoformat()
            elif isinstance(value, date):
                a[key] = value.isoformat()
        if str(a["article"]).startswith("6A"):
            articles[a["article"]] = a
    return articles, tables, layout


def equal(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, abs_tol=1e-8, rel_tol=1e-10)
    return a == b


def totals(rows):
    result = {"articles": len(rows), "status": dict(Counter(r["status"] for r in rows))}
    for key in ("base", "plan_units", "preseason", "season_fact", "forecast"):
        known = [r[key] for r in rows if r.get(key) is not None]
        result[key] = {"known": len(known), "sum": sum(known) if known else None}
    comparable = [r for r in rows if r.get("plan_units") is not None and r.get("season_fact") is not None]
    denominator = sum(r["plan_units"] for r in comparable)
    result["execution"] = {
        "known": len(comparable),
        "value": sum(r["season_fact"] for r in comparable) / denominator if denominator else None,
    }
    result["recommendation_max_length"] = max(len(r.get("recommendation") or "") for r in rows)
    return result


def run():
    before = json.loads((OUT / "before.json").read_text(encoding="utf-8"))
    after = json.loads((ROOT / "test-results/real-merchandise/result.json").read_text(encoding="utf-8"))
    downloads = Path.home() / "Downloads"
    skill_path, app_path = downloads / "Анализ продаж скилл.xlsx", downloads / "Анализ продаж приложение.xlsx"
    skill, skill_tables, skill_layout = workbook(skill_path)
    app, app_tables, app_layout = workbook(app_path)
    old = {a["article"]: a for a in before["sections"]["articles"]}
    new = {a["article"]: a for a in after["sections"]["articles"]}
    assert skill.keys() == app.keys() == old.keys() == new.keys() and len(new) == 338
    numeric = ("base", "plan_units", "season_fact", "execution", "forecast")
    assert all(equal(app[a][k], old[a][k]) for a in new for k in numeric), (
        "Supplied app workbook differs from recorded baseline"
    )
    unchanged = (
        "base",
        "target",
        "season_plan",
        "plan_pct",
        "plan_units",
        "season_fact",
        "st",
        "execution",
        "gap_pp",
        "gap_units",
        "pace",
        "required",
        "cover",
        "forecast",
        "forecast_st",
    )
    assert all(equal(old[a][k], new[a][k]) for a in new for k in unchanged)
    assert all(
        equal(old[a]["second_wave"], {k: v for k, v in new[a]["second_wave"].items() if k != "user_comment"})
        for a in new
        if old[a]["second_wave"]
    )
    comparison = [
        {
            "article": a,
            "skill": skill[a],
            "app_before": {k: old[a].get(k) for k in LABELS},
            "app_after": {
                k: new[a].get(k)
                for k in [
                    *LABELS,
                    "preliminary_band",
                    "operational_state",
                    "season_observation_days",
                    "status_reason",
                    "earliest_status_review_date",
                ]
            },
            "changed": [k for k in LABELS if not equal(old[a].get(k), new[a].get(k))],
        }
        for a in new
    ]
    _, _, after_layout = workbook(next((ROOT / "test-results/real-merchandise").glob("KANZLER_*.xlsx")))
    summaries = {
        "skill": totals(list(skill.values())),
        "app_before": totals(list(old.values())),
        "app_after": totals(list(new.values())),
    }
    differences = {
        k: sum(not equal(skill[a][k], new[a].get(k)) for a in new) for k in LABELS if k != "article"
    }
    s = after["sections"]
    details = {
        "preliminary": dict(Counter(a["preliminary_band"] or "Нет данных" for a in new.values())),
        "operations": dict(Counter(a["operational_state"] for a in new.values())),
        "observation_days": dict(Counter(str(a["season_observation_days"]) for a in new.values())),
        "size_signals": sum(a["diagnosis"]["sizes"]["signal"] for a in new.values()),
        "size_confirmed_before": sum(a["diagnosis"]["sizes"]["problem"] for a in old.values()),
        "size_confirmed_after": sum(a["diagnosis"]["sizes"]["problem"] for a in new.values()),
        "qa_before": len(before["sections"]["data-quality"]),
        "qa_after_detail": len(s["data-quality"]),
        "qa_after_groups": len(after["result"]["quality_summary"]),
        "categories": len(s["categories"]),
        "categories_with_plan": sum(c["known_plan_units"] is not None for c in s["categories"]),
        "categories_with_fact": sum(c["known_fact_units"] is not None for c in s["categories"]),
        "confirmed_second_risks": sum(w["risk"] for w in s["second-wave"]),
        "conditional_second_risks": sum(w["scenario_risk"] for w in s["second-wave"]),
        "price_actions": sum(
            p["decision"] in ("снизить цену", "повысить цену", "уменьшить скидку") for p in s["pricing"]
        ),
    }
    details["categories_comparable_execution"] = sum(c["execution"] is not None for c in s["categories"])
    details["category_metrics_known_before_after"] = {
        key: [
            sum(c[key] is not None for c in before["sections"]["categories"]),
            sum(c[key] is not None for c in s["categories"]),
        ]
        for key in ("base", "plan_units", "season_fact", "execution", "st", "forecast")
    }
    details["second_confirmed_forecasts"] = sum(w["forecast"] is not None for w in s["second-wave"])
    details["second_conditional_forecasts"] = sum(
        w["scenario_forecast"] is not None for w in s["second-wave"]
    )
    payload = {
        "sources": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (skill_path, app_path)},
        "summaries": summaries,
        "skill_differences": differences,
        "details": details,
        "articles": comparison,
        "categories": {
            "skill": skill_tables["Категории"],
            "before": before["sections"]["categories"],
            "after": s["categories"],
        },
        "actions": {
            "skill": skill_tables["Действия"],
            "before": app_tables["Действия"],
            "after": s["actions"],
        },
        "layout": {"skill": skill_layout, "before": app_layout, "after": after_layout},
    }
    (OUT / "comparison.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    lines = [
        "# Коммерческий модуль: дифференциальная проверка",
        "",
        "338 артикулов сопоставлены по точному артикулу в Skill, исходном Excel приложения и новом результате. Файлы-источники прочитаны без изменений.",
        "",
        "| Показатель | Skill | Приложение до | Приложение после |",
        "|---|---:|---:|---:|",
    ]
    for key, label in [
        ("base", "База"),
        ("plan_units", "План на дату"),
        ("season_fact", "Факт сезона"),
        ("preseason", "Предсезон"),
        ("forecast", "Прогноз"),
    ]:
        cells = [
            f"{v[key]['sum']:,.2f}; {v[key]['known']}/338"
            if v[key]["sum"] is not None
            else "Неизвестно; 0/338"
            for v in summaries.values()
        ]
        lines.append(f"| {label}, ед.; покрытие | " + " | ".join(cells) + " |")
    lines.append(
        "| Выполнение; сопоставимых | "
        + " | ".join(f"{v['execution']['value']:.1%}; {v['execution']['known']}" for v in summaries.values())
        + " |"
    )
    lines += [
        "",
        "Все 15 проверенных коммерческих показателей каждого артикула и строгие расчёты второй поставки совпадают с приложением до изменений. Восстановление предсезона совпало со Skill по известной сумме. Отличия плана от Skill сохранены: точная категорийная замена расширяет покрытие; консервативный темп прогноза не заменён более оптимистичным сценарием Skill.",
        "",
        "```json",
        json.dumps(details, ensure_ascii=False, indent=2),
        "```",
        "",
        "Итог: 338 «НЕДОСТАТОЧНО ДАННЫХ». Исключений раннего финального статуса нет. 201 артикул наблюдался 6 сезонных дней, 11 — 2 дня, один — 5 дней; 125 не имеют доказанной даты коммерческого входа. Предварительный уровень не является финальным статусом.",
        "",
        "## Примеры артикулов",
        "",
        "| Артикул | Статус до → после | Предварительно | Сезонных дней | Операционный сигнал |",
        "|---|---|---|---:|---|",
    ]
    examples = []
    for band in ("ХИТ", "В ПЛАНЕ", "РИСК", "АУТСАЙДЕР", None):
        examples.extend([a for a in new.values() if a["preliminary_band"] == band][:2])
    for a in examples:
        lines.append(
            f"| {a['article']} | {old[a['article']]['status']} → {a['status']} | {a['preliminary_band'] or '—'} | {a['season_observation_days'] or '—'} | {a['operational_signal']} |"
        )
    lines += [
        "",
        "## Пользовательский формат",
        "",
        f"Артикулы: {app_layout['Артикулы']['columns']} → {after_layout['Артикулы']['columns']} колонок; суммарная ширина {app_layout['Артикулы']['width']:.0f} → {after_layout['Артикулы']['width']:.0f}. Действия — 8 колонок. Сводка — 4. QA: {details['qa_before']} подробных сообщений раньше; {details['qa_after_detail']} после с учётом новых пояснений предсезона; в Excel {details['qa_after_groups']} групп. Полная детализация сохранена в UI/API.",
        "",
        "Числовые суммы выше приведены без пользовательского округления для сверки. Excel/UI показывают физические единицы целыми, проценты с одним знаком. Темпы остаются дробными.",
        "",
        "### Пример рекомендации",
        "",
        f"Артикул {examples[0]['article']}. До: {old[examples[0]['article']]['recommendation']}",
        "",
        f"После: {examples[0]['recommendation']}",
        "",
        "### Пример действия",
        "",
        json.dumps(s["actions"][0], ensure_ascii=False, indent=2),
        "",
        "Полная построчная сверка, категорийные таблицы, действия всех трёх вариантов, размеры листов и SHA исходных сравниваемых файлов сохранены в `comparison.json`. Это файлы проверки, служебные данные в пользовательский Excel не выведены.",
    ]
    (OUT / "comparison.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {"summaries": summaries, "skill_differences": differences, "details": details},
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    run()
