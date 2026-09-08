from datetime import date

from app.core.exceptions import IntegrityError
from app.modules.merchandise.presentation import UNIT_KEYS, display_units, group_quality
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Shared presentation schema: web and XLSX render the same backend result.
ARTICLE_COLUMNS = [
    ("article", "Артикул"),
    ("code", "Код"),
    ("category", "Категория"),
    ("kind", "Вид номенклатуры"),
    ("assortment", "Вид ассортимента"),
    ("strategy_label", "Стратегия оценки"),
    ("entry", "Дата входа"),
    ("season_observation_days", "Сезонное наблюдение, дней"),
    ("base", "Начальный остаток"),
    ("target", "Целевой процент реализации сезона, %"),
    ("season_plan", "План сезона, ед."),
    ("plan_pct", "План на текущую дату, % накоп."),
    ("plan_units", "План на текущую дату, ед."),
    ("preseason", "Продажи до начала сезона, ед. — оценка"),
    ("season_fact", "Факт продаж сезона, ед."),
    ("st", "Факт реализации сезона, %"),
    ("execution", "Выполнение плана на текущую дату, %"),
    ("gap_pp", "Отклонение, п.п."),
    ("gap_units", "Отклонение, ед."),
    ("previous_week", "Продажи предыдущей недели, ед."),
    ("current_week", "Продажи текущей недели, ед."),
    ("avg2", "Среднее за 2 недели, ед./нед."),
    ("avg4", "Среднее за 3–4 недели, ед./нед."),
    ("required", "Требуемый темп, ед./нед."),
    ("pace_ratio", "Фактический темп / требуемый, %"),
    ("cover", "Покрытие запасом, недель"),
    ("forecast", "Прогноз продаж, ед."),
    ("forecast_st", "Прогноз реализации, %"),
    ("preliminary_band", "Предварительный уровень выполнения"),
    ("status", "Коммерческий статус"),
    ("operational_state", "Операционный статус"),
    ("operational_signal", "Операционный сигнал"),
    ("primary_cause", "Основная причина"),
    ("recommendation", "Рекомендация"),
    ("owner", "Ответственный"),
    ("review_date", "Контрольная дата"),
    ("review_window", "Окно проверки без утверждённой даты"),
    ("stock", "Остаток"),
]
CATEGORY_COLUMNS = [
    ("category", "Категория"),
    ("assortment", "Вид ассортимента"),
    ("known_base_units", "Известный начальный остаток, ед."),
    ("base_known_articles", "Артикулов с базой"),
    ("plan_pct", "План на дату, %"),
    ("known_plan_units", "Известный план на дату, ед."),
    ("plan_known_articles", "Артикулов с планом"),
    ("known_fact_units", "Известный факт сезона, ед."),
    ("fact_known_articles", "Артикулов с фактом"),
    ("st", "Факт реализации, %"),
    ("st_comparable_articles", "Сопоставимых для реализации"),
    ("execution", "Выполнение плана, %"),
    ("execution_comparable_articles", "Сопоставимых для выполнения"),
    ("previous_week", "Предыдущая неделя, ед."),
    ("current_week", "Текущая неделя, ед."),
    ("avg2", "Среднее за 2 недели, ед./нед."),
    ("known_forecast_units", "Известный прогноз, ед."),
    ("forecast_known_articles", "Артикулов с прогнозом"),
    ("hits", "Хиты"),
    ("on_plan", "В плане"),
    ("risks", "Риск"),
    ("outsiders", "Аутсайдер"),
    ("total_articles", "Всего артикулов"),
    ("conclusion", "Вывод"),
]
PRICING_COLUMNS = [
    ("article", "Артикул"),
    ("status", "Статус"),
    ("price", "Текущая цена"),
    ("discount", "Текущая скидка"),
    ("markup", "Наценка"),
    ("margin", "Маржа"),
    ("decision", "Решение"),
    ("change", "Предлагаемое изменение"),
    ("reason", "Причина"),
    ("review_date", "Контрольная дата"),
    ("review_window", "Окно проверки"),
]
SECOND_COLUMNS = [
    ("article", "Артикул"),
    ("base", "Начальный остаток"),
    ("first_sales", "Факт продаж до 2-й поставки, ед."),
    ("first_st", "Реализация до 2-й поставки, %"),
    ("target70", "Цель 70%, ед."),
    ("remaining70", "Остаток до цели 70%, ед."),
    ("second_date", "Дата 2-й поставки"),
    ("second_qty", "Объём 2-й поставки"),
    ("weeks", "Недель до поставки"),
    ("completed_weeks", "Завершённых недель в оценке темпа"),
    ("required", "Требуемый темп, ед./нед."),
    ("pace", "Текущий темп, ед./нед."),
    ("pace_ratio", "Фактический темп / требуемый, %"),
    ("forecast", "Прогноз к дате поставки, ед."),
    ("forecast_st", "Прогноз реализации, %"),
    ("gap", "Отклонение"),
    ("decision", "Решение"),
    ("user_comment", "Комментарий"),
    ("scenario_forecast", "Условный прогноз без подтверждения партии, ед."),
    ("scenario_risk", "Условный недобор без подтверждения партии"),
]
ACTION_COLUMNS = [
    ("priority", "Приоритет"),
    ("article", "Артикул / категория"),
    ("problem", "Проблема"),
    ("action", "Действие"),
    ("owner", "Ответственный"),
    ("review_date", "Контрольная дата"),
    ("expected", "Ожидаемый результат"),
    ("review_window", "Окно проверки"),
]
QA_COLUMNS = [
    ("severity", "Уровень"),
    ("article", "Артикул"),
    ("source", "Источник"),
    ("source_row", "Строка источника"),
    ("message", "Комментарий"),
]
PERCENT_KEYS = {
    "target",
    "plan_pct",
    "st",
    "execution",
    "wow",
    "pace_ratio",
    "forecast_st",
    "size_availability",
    "broken_ratio",
    "warehouse_share",
    "distribution",
    "discount",
    "margin",
    "first_st",
}
SCHEMAS = {
    "articles": ARTICLE_COLUMNS,
    "categories": CATEGORY_COLUMNS,
    "pricing": PRICING_COLUMNS,
    "second-wave": SECOND_COLUMNS,
    "actions": ACTION_COLUMNS,
    "data-quality": QA_COLUMNS,
    "monthly-plan": [
        ("article", "Артикул"),
        ("category", "Категория"),
        ("base", "Начальный остаток"),
        ("target", "Цель реализации"),
        ("month", "Месяц"),
        ("plan_pct", "Доля месяца от базы"),
        ("monthly_units", "План месяца, ед."),
    ],
    "history": [("article", "Артикул"), ("category", "Категория"), ("context", "Исторический контекст")],
    "methodology": [("topic", "Правило"), ("rule", "Методика и ограничения")],
}
SHEETS = {
    "summary": "Сводка",
    "categories": "Категории",
    "articles": "Артикулы",
    "weekly-plan": "План по неделям",
    "pricing": "Ценовые решения",
    "second-wave": "Вторая волна",
    "actions": "Действия",
    "data-quality": "Качество данных",
    "monthly-plan": "План по месяцам",
    "history": "История",
    "methodology": "Методика",
}
SUMMARY_COLUMNS = [
    ("label", "Показатель"),
    ("value", "Значение"),
    ("coverage", "Покрытие"),
    ("meaning", "Управленческий смысл"),
]
QUALITY_GROUP_COLUMNS = [
    ("severity", "Уровень"),
    ("source", "Источник"),
    ("message", "Комментарий"),
    ("affected_articles", "Затронуто артикулов"),
    ("occurrences", "Сообщений"),
]


def export(result, sections, path):
    wb = Workbook()
    wb.remove(wb.active)
    summary = [{"label": k, "value": v} for k, v in result["display_summary"].items()]
    for section, title in SHEETS.items():
        ws = wb.create_sheet(title)
        if section == "summary":
            columns = (
                SUMMARY_COLUMNS
                if result.get("summary_rows")
                else [("label", "Показатель"), ("value", "Значение")]
            )
            rows = result.get("summary_rows", summary)
        elif section == "data-quality":
            columns, rows = QUALITY_GROUP_COLUMNS, group_quality(sections[section])
        elif section == "weekly-plan":
            columns = [
                ("article", "Артикул"),
                ("code", "Код"),
                ("base", "Начальный остаток"),
                ("target", "Целевой процент реализации сезона, %"),
            ]
            columns += [(point["week_start"], point["label"]) for point in result["week_columns"]]
            rows = sections[section]
        else:
            columns = SCHEMAS[section]
            rows = sections[section]
        ws.append([label for key, label in columns])
        for row in rows:
            values = []
            for key, label in columns:
                value = row.get(key)
                if key in UNIT_KEYS or (
                    section == "summary" and key == "value" and row.get("format") == "units"
                ):
                    value = display_units(value)
                if key in ("entry", "review_date", "second_date") and value:
                    value = date.fromisoformat(value)
                values.append(value)
            ws.append(values)
        ws.freeze_panes = "D2" if section not in ("summary", "data-quality") else "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.zoomScale = 85
        ws.row_dimensions[1].height = 42
        for ci, (key, label) in enumerate(columns, 1):
            width = 12
            if key in ("article",):
                width = 18
            elif key in (
                "category",
                "kind",
                "assortment",
                "type",
                "status",
                "decision",
                "preliminary_band",
                "operational_state",
                "coverage",
                "value",
            ):
                width = 20
            elif key in ("primary_cause", "secondary_cause", "reason", "owner"):
                width = 26
            elif key in (
                "recommendation",
                "action",
                "comment",
                "qa_text",
                "historical_text",
                "message",
                "expected",
                "change",
                "conclusion",
                "label",
                "norms_text",
                "economics_text",
                "evidence_text",
                "missing_evidence",
                "review_limitation",
                "operational_signal",
                "meaning",
                "rule",
                "context",
                "user_comment",
                "problem",
            ):
                width = 42
            ws.column_dimensions[get_column_letter(ci)].width = width
            if section == "methodology":
                ws.column_dimensions[get_column_letter(ci)].width = 80 if key == "rule" else 20
            for ri in range(1, ws.max_row + 1):
                c = ws.cell(ri, ci)
                c.font = Font(name="Calibri", size=9, bold=ri == 1, color="FFFFFF" if ri == 1 else "172B4D")
                c.alignment = Alignment(
                    vertical="center",
                    horizontal="center"
                    if ri == 1
                    else "right"
                    if isinstance(c.value, (int, float))
                    else "left",
                    wrap_text=True,
                )
                if ri == 1:
                    c.fill = PatternFill("solid", fgColor="24384B")
                elif ri % 2 == 0:
                    c.fill = PatternFill("solid", fgColor="F2F5F8")
                if ri > 1:
                    c.number_format = (
                        "0.0%"
                        if key in PERCENT_KEYS
                        or (section == "weekly-plan" and ci >= 5)
                        or (
                            section == "summary"
                            and key == "value"
                            and rows[ri - 2].get("format") == "percent"
                        )
                        else "dd.mm.yyyy"
                        if isinstance(c.value, date)
                        else "#,##0"
                        if key in UNIT_KEYS
                        or (section == "summary" and key == "value" and rows[ri - 2].get("format") == "units")
                        else "#,##0.0"
                        if isinstance(c.value, (int, float))
                        else "@"
                    )
                    if isinstance(c.value, str) and c.value.startswith(("=", "+", "-", "@")):
                        c.data_type = "s"
            for ri in range(2, ws.max_row + 1):
                ws.row_dimensions[ri].height = 36 if section in ("actions", "data-quality", "pricing") else 22
                if section == "methodology":
                    ws.row_dimensions[ri].height = 48
    wb.save(path)
    check = load_workbook(path, read_only=False, data_only=True)
    if check.sheetnames != list(SHEETS.values()):
        raise IntegrityError("Не совпадает состав листов отчёта.")
    forbidden = re_forbidden()
    for ws in check:
        if not ws.freeze_panes or not ws.auto_filter.ref:
            raise IntegrityError("Не сохранены фильтры и закрепление областей.")
        if any(forbidden.search(str(c.value or "")) for c in ws[1]):
            raise IntegrityError("В пользовательский отчёт попал технический заголовок.")
    if check["Артикулы"].max_row != len(sections["articles"]) + 1:
        raise IntegrityError("Количество артикулов в Excel не совпало с результатом анализа.")
    check.close()


def re_forbidden():
    import re

    return re.compile(
        r"\b(PLAN|FACT|FORECAST|HIST|ST|Gap|Stock Cover|Markup|SKU|Avg)\b|1-я партия|Строка FACT",
        re.IGNORECASE,
    )
