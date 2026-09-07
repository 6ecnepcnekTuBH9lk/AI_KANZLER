import json
from pathlib import Path

from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
downloads = Path("C:/Users/Ermolenko.i/Downloads")
names = [
    "Анализ продаж (Общий) (XLSX) 2026-09-06.xlsx",
    "План сезона ОЗ 26-27 из ценообразования.xlsx",
    "ОЗ26_W33_план_сезона 28.08+фото.xlsx",
    "Артикулы_коллеции_ВЛ_27_Нет_кодов_ТНВЭД_2.xlsx",
]
for i, name in enumerate(names):
    wb = load_workbook(downloads / name, read_only=True, data_only=True)
    result = []
    for ws in wb:
        sample = []
        for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
            if ri > 15:
                break
            sample.append(
                {"row": ri, "cells": {str(c + 1): str(v) for c, v in enumerate(row) if v is not None}}
            )
        result.append({"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column, "sample": sample})
    (root / f"docs/source-review/input-{i}.json").write_text(
        json.dumps({"file": name, "sheets": result}, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(i, [(s["sheet"], s["rows"], s["columns"]) for s in result])
    wb.close()
