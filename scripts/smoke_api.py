"""Exercise the running local app with actual inputs and download both results."""

import hashlib
import io
import json
import time
from contextlib import ExitStack
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
source = Path.home() / "Downloads"
names = [
    "ОЗ26_W33_план_сезона 28.08+фото.xlsx",
    "План сезона ОЗ 26-27 из ценообразования.xlsx",
    "Анализ продаж (Общий) (XLSX) 2026-09-06.xlsx",
]
with httpx.Client(base_url="http://127.0.0.1:8000", timeout=120) as client:
    assert client.get("/api/health").json()["status"] == "ok"
    with ExitStack() as stack:
        files = [("files", (n, stack.enter_context((source / n).open("rb")))) for n in names]
        response = client.post("/api/merchandise/runs", files=files)
    assert response.status_code == 202, response.text
    identifier = response.json()["id"]
    print("Коммерческий запуск:", identifier, flush=True)
    for _ in range(240):
        job = client.get(f"/api/jobs/{identifier}").json()
        if job["status"] in ("completed", "failed"):
            break
        time.sleep(1)
    assert job["status"] == "completed", job
    assert job["result"]["summary"]["count"] == 338
    assert job["created_at"].endswith("+00:00")
    for file in job["files"]:
        assert hashlib.sha256((source / file["name"]).read_bytes()).hexdigest() == file["sha256"]
    history = client.get("/api/history").json()["items"]
    tnved = next(j for j in history if j["module"] == "tnved" and j["status"] == "completed")
    output = ROOT / "test-results/api-smoke"
    output.mkdir(parents=True, exist_ok=True)
    for j in (job, tnved):
        response = client.get(f"/api/{j['module']}/runs/{j['id']}/export")
        assert response.status_code == 200
        wb = load_workbook(io.BytesIO(response.content), read_only=True)
        assert len(wb.sheetnames) == 8 if j["module"] == "merchandise" else len(wb.sheetnames) >= 1
        wb.close()
        (output / j["output_name"]).write_bytes(response.content)
    (output / "jobs.json").write_text(
        json.dumps([job, tnved], ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("Проверены 338 артикулов, история после рестарта, SHA256 входов и два скачанных XLSX.")
