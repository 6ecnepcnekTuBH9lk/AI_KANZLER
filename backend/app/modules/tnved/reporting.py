"""Patch only target OOXML cells; untouched ZIP members remain byte-identical."""

from copy import deepcopy
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

from app.core.exceptions import IntegrityError
from lxml import etree
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def tag(name):
    return "{" + NS + "}" + name


def patch_workbook(source, destination, edits, added_columns):
    """edits: {sheet: {coordinate: (text, code_as_text)}}."""
    if Path(source).resolve() == Path(destination).resolve():
        raise IntegrityError("Исходный Excel нельзя перезаписывать.")
    with ZipFile(source) as zin:
        parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
        workbook = etree.fromstring(parts["xl/workbook.xml"])
        relations = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
        rels = {r.get("Id"): r.get("Target") for r in relations}
        sheet_paths = {}
        for s in workbook.find(tag("sheets")):
            target = rels[s.get("{" + REL + "}id")]
            sheet_paths[s.get("name")] = (
                target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
            )
        styles = etree.fromstring(parts["xl/styles.xml"])
        xfs = styles.find(tag("cellXfs"))
        new_styles = {}

        def style_id(base, code=False, wrap=False):
            key = (base, code, wrap)
            if key not in new_styles:
                xf = deepcopy(xfs[int(base)])
                if code:
                    xf.set("numFmtId", "49")
                    xf.set("applyNumberFormat", "1")
                if wrap:
                    alignment = xf.find(tag("alignment"))
                    if alignment is None:
                        alignment = etree.SubElement(xf, tag("alignment"))
                    alignment.set("wrapText", "1")
                    xf.set("applyAlignment", "1")
                new_styles[key] = str(len(xfs))
                xfs.append(xf)
            return new_styles[key]

        changed = set()
        for sheet, changes in edits.items():
            path = sheet_paths[sheet]
            root = etree.fromstring(parts[path])
            data = root.find(tag("sheetData"))
            rows = {int(r.get("r")): r for r in data}
            for coord, (value, is_code) in changes.items():
                col, ri = coordinate_from_string(coord)
                ci = column_index_from_string(col)
                row = rows.get(ri)
                if row is None:
                    raise IntegrityError("Нельзя добавлять товарные строки при заполнении кодов.")
                cells = {c.get("r"): c for c in row}
                cell = cells.get(coord)
                if cell is not None and cell.find(tag("f")) is not None:
                    raise IntegrityError("Нельзя перезаписывать существующую формулу Excel.")
                new = cell is None
                if new:
                    cell = etree.Element(tag("c"), r=coord)
                    neighbor = cells.get(f"{get_column_letter(ci - 1)}{ri}") if ci > 1 else None
                    if neighbor is not None and neighbor.get("s"):
                        cell.set("s", neighbor.get("s"))
                    at = next(
                        (
                            i
                            for i, c in enumerate(row)
                            if column_index_from_string(coordinate_from_string(c.get("r"))[0]) > ci
                        ),
                        len(row),
                    )
                    row.insert(at, cell)
                cell.set("s", style_id(cell.get("s", "0"), code=is_code, wrap=new and not is_code))
                for child in list(cell):
                    if child.tag in (tag("f"), tag("v"), tag("is")):
                        cell.remove(child)
                cell.set("t", "inlineStr")
                inline = etree.SubElement(cell, tag("is"))
                text = etree.SubElement(inline, tag("t"))
                text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                text.text = "" if value is None else str(value)
            extra = added_columns.get(sheet, {})
            if extra:
                cols = root.find(tag("cols"))
                if cols is None:
                    cols = etree.Element(tag("cols"))
                    root.insert(list(root).index(data), cols)
                for ci, width in extra.items():
                    # Reuse an exact column definition on repeat exports.
                    existing_col = next(
                        (c for c in cols if c.get("min") == str(ci) and c.get("max") == str(ci)), None
                    )
                    if existing_col is not None:
                        existing_col.set("width", str(width))
                        existing_col.set("customWidth", "1")
                        continue
                    etree.SubElement(
                        cols, tag("col"), min=str(ci), max=str(ci), width=str(width), customWidth="1"
                    )
                dim = root.find(tag("dimension"))
                if dim is not None:
                    old_col, old_row = coordinate_from_string(dim.get("ref").split(":")[-1])
                    last_col = max(max(extra), column_index_from_string(old_col))
                    dim.set("ref", f"A1:{get_column_letter(last_col)}{max(max(rows), old_row)}")
            parts[path] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            changed.add(path)
        if new_styles:
            xfs.set("count", str(len(xfs)))
            parts["xl/styles.xml"] = etree.tostring(
                styles, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            changed.add("xl/styles.xml")
        with ZipFile(destination, "w") as zout:
            zout.comment = zin.comment
            for info in zin.infolist():
                zout.writestr(info, parts[info.filename])
    # Check the archive actually saved and all non-target parts survived exactly.
    with ZipFile(source) as original, ZipFile(destination) as result:
        if original.namelist() != result.namelist():
            raise IntegrityError("Изменился состав исходного Excel.")
        for name in original.namelist():
            if name not in changed and original.read(name) != result.read(name):
                raise IntegrityError("Изменился служебный компонент исходного Excel.")
        for sheet, changes in edits.items():
            path = sheet_paths[sheet]
            before, after = etree.fromstring(original.read(path)), etree.fromstring(result.read(path))
            before_cells = {c.get("r"): c for c in before.iter(tag("c"))}
            after_cells = {c.get("r"): c for c in after.iter(tag("c"))}
            if [r.get("r") for r in before.iter(tag("row"))] != [r.get("r") for r in after.iter(tag("row"))]:
                raise IntegrityError("Изменился порядок строк исходного Excel.")
            for coord, c in before_cells.items():
                if coord not in changes and etree.tostring(c) != etree.tostring(after_cells.get(coord)):
                    raise IntegrityError("Изменилась исходная ячейка вне колонок результата.")
