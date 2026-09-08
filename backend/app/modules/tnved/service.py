from dataclasses import asdict
from pathlib import Path

from app.core.exceptions import InputError, IntegrityError, SourceUnavailable
from app.modules.tnved.alta import AltaClient
from app.modules.tnved.classifier import ClassificationEngine, valid_code
from app.modules.tnved.comments import user_comment
from app.modules.tnved.features import extract_features
from app.modules.tnved.reporting import patch_workbook
from app.shared.excel import norm, read_tables
from openpyxl.utils import get_column_letter


def run(files, options, output_dir, progress, sessions, client_factory=AltaClient):
    if len(files) != 1:
        raise InputError("Загрузите один XLSX для классификации.")
    file = files[0]
    tables = read_tables(file.path)
    formula_cells = {
        (table.sheet, ri, i): cell.value
        for table in read_tables(file.path, data_only=False)
        for ri, row in table.rows
        for i, cell in enumerate(row)
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }
    items, edits, added = [], {}, {}
    progress(5, "Чтение номенклатуры")
    client = client_factory()
    engine = ClassificationEngine(client, sessions)
    try:
        total = sum(len(t.rows) for t in tables)
        for table in tables:
            if table.col("номенклатура", "наименование") is None:
                continue
            code_cols = [i for i, h in enumerate(table.headers) if re_code_header(h)]
            if len(code_cols) > 1:
                raise InputError(f"На листе «{table.sheet}» несколько колонок кода ТН ВЭД.")
            # Ignore styled empty trailing columns, but never reuse a data/formula column.
            occupied = {
                i + 1
                for row in table.preamble + [row for _, row in table.rows]
                for i, cell in enumerate(row)
                if cell.value not in (None, "")
            }
            occupied.update(ci + 1 for sheet, ri, ci in formula_cells if sheet == table.sheet)
            code_col = code_cols[0] + 1 if code_cols else max(occupied, default=0) + 1
            occupied.add(code_col)
            comments_col = table.col("Комментарий ТН ВЭД")
            if comments_col is None:
                comments_col = table.col("Комментарий")
            comments_col = comments_col + 1 if comments_col is not None else None
            formula_comments = comments_col is not None and any(
                sheet == table.sheet and ci == comments_col - 1 for sheet, ri, ci in formula_cells
            )
            if formula_comments:
                comments_col = None
            if comments_col is None:
                comments_col = max(occupied) + 1
            occupied.add(comments_col)
            source_cols = table.columns(("Источник", "Источник Alta.ru"))
            source_col = next(
                (
                    i + 1
                    for i in source_cols
                    if not any(sheet == table.sheet and ci == i for sheet, ri, ci in formula_cells)
                ),
                max(occupied) + 1,
            )
            sheet_edits = edits.setdefault(table.sheet, {})
            sheet_added = added.setdefault(table.sheet, {})
            for col, header, width in (
                (code_col, "Код ТНВЭД", 16),
                (comments_col, "Комментарий ТН ВЭД" if formula_comments else "Комментарий", 60),
                (source_col, "Источник", 54),
            ):
                # Preserve the dedicated header when reprocessing a formula workbook.
                if col == comments_col and table.col("Комментарий ТН ВЭД") == col - 1:
                    header = "Комментарий ТН ВЭД"
                sheet_edits[f"{get_column_letter(col)}{table.header_row}"] = (header, False)
                sheet_added[col] = width
            for ri, row in table.rows:
                name = table.get(row, "Наименование") or table.get(row, "Номенклатура")
                if not norm(name) or norm(name) in ("итого", "всего", "общий итог"):
                    continue
                existing = row[code_col - 1].value if code_cols else None
                cached_value = existing
                existing = formula_cells.get((table.sheet, ri, code_col - 1), existing)
                data = {h: row[i].value for i, h in enumerate(table.headers) if h}
                features = extract_features(data)
                article = str(table.get(row, "Артикул", "Артикулы") or table.get(row, "Код") or "")
                item = {
                    "id": f"{table.sheet}:{ri}",
                    "sheet": table.sheet,
                    "row": ri,
                    "article": article,
                    "source_code": str(table.get(row, "Код") or ""),
                    "nomenclature": str(table.get(row, "Номенклатура") or ""),
                    "source_name": str(table.get(row, "Наименование") or ""),
                    "name": str(name),
                    "original_code": existing,
                    "signature": features.signature,
                    "features": asdict(features),
                }
                if existing not in (None, "") and not options.get("verify_existing", False):
                    item.update(
                        code=existing,
                        status="Уже имел код",
                        comment="Исходный код сохранён без проверки.",
                        evidence=None,
                        candidates=[],
                    )
                else:
                    result = engine.classify(features)
                    item.update(result)
                    if result["code"] and (not valid_code(result["code"]) or not result["evidence"]):
                        raise IntegrityError("Нельзя записать непроверенный код ТН ВЭД.")
                    is_formula = (table.sheet, ri, code_col - 1) in formula_cells
                    if existing not in (None, ""):
                        value = cached_value if is_formula else existing
                        value = str(value) if value is not None else None
                        item["existing_verification"] = {"status": "UNPROVEN", "evidence": None}
                        if valid_code(value):
                            try:
                                proof = client.verify(value)
                                verified = result["code"] == value
                                item["existing_verification"] = {
                                    "status": "CONFIRMED" if verified else "UNPROVEN",
                                    "evidence": proof,
                                }
                            except SourceUnavailable as exc:
                                item["existing_verification"]["reason"] = exc.message
                        else:
                            item["existing_verification"]["reason"] = (
                                "Нет вычисленного значения формулы для проверки."
                                if is_formula
                                else "Исходное значение не является кодом из 10 цифр."
                            )
                        if is_formula or not result["code"]:
                            item["proposed_code"] = result["code"]
                            if item["existing_verification"]["status"] != "CONFIRMED":
                                item.update(
                                    code=None,
                                    evidence=None,
                                    status="Исходный код не подтверждён",
                                    comment="Исходная формула сохранена. "
                                    if is_formula
                                    else "Исходный код сохранён, но не подтверждён. ",
                                )
                                item["comment"] += (
                                    item["existing_verification"].get("reason", "") + " " + result["comment"]
                                )
                    if not is_formula and (existing in (None, "") or result["code"]):
                        sheet_edits[f"{get_column_letter(code_col)}{ri}"] = (result["code"] or "", True)
                item["original_comment"] = row[comments_col - 1].value if comments_col <= len(row) else None
                item["user_comment"] = user_comment(features, item)
                item["source_url"] = (
                    f"https://www.alta.ru/tnved/code/{item['code']}/"
                    if item["status"] == "Код определён" and valid_code(item["code"]) and item["evidence"]
                    else ""
                )
                sheet_edits[f"{get_column_letter(comments_col)}{ri}"] = (item["user_comment"], False)
                sheet_edits[f"{get_column_letter(source_col)}{ri}"] = (item["source_url"], False)
                items.append(item)
                progress(
                    min(90, 5 + int(len(items) / max(total, 1) * 85)), f"Обработано товаров: {len(items)}"
                )
    finally:
        client.close()
    if not items:
        raise InputError("В Excel не найдены товарные строки с наименованием.")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_name = Path(file.original_name).stem + "_с_кодами_ТНВЭД.xlsx"
    path = output_dir / output_name
    progress(95, "Проверка сохранности Excel")
    patch_workbook(file.path, path, edits, added)
    counts = {
        s: sum(i["status"] == s for i in items)
        for s in (
            "Уже имел код",
            "Код определён",
            "Требуется уточнение",
            "Техническая ошибка",
            "Исходный код не подтверждён",
        )
    }
    return (
        {"title": file.original_name, "total": len(items), "counts": counts},
        {"items": items},
        path,
        output_name,
        [],
    )


def re_code_header(value):
    return norm(value).replace(" ", "") == "кодтнвэд"
