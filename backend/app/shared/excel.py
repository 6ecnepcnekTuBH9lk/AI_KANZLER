"""Semantic XLSX reader. Source cells retain type and number format."""

import math
import re
from dataclasses import dataclass
from datetime import date, datetime

from app.core.exceptions import InputError
from openpyxl import load_workbook


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ").replace("ё", "е").strip().lower())


def number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        n = float(str(value).replace("\xa0", "").replace(" ", "").replace(",", ".").rstrip("%"))
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None


def fraction(value, fmt="", whole=False):
    n = number(value)
    if n is None:
        return None
    if "%" in str(value) or (whole and "%" not in fmt) or n > 1:
        return n / 100
    return n


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value or "").strip()
    for pattern, fmt in [(r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d"), (r"\d{2}\.\d{2}\.\d{4}", "%d.%m.%Y")]:
        match = re.search(pattern, s)
        if match:
            try:
                # Calendar dates intentionally carry no timezone.
                return date.fromisoformat(
                    match[0] if fmt == "%Y-%m-%d" else "-".join(reversed(match[0].split(".")))
                )
            except ValueError:
                return None
    return None


@dataclass
class Cell:
    value: object
    fmt: str = ""


@dataclass
class Table:
    sheet: str
    header_row: int
    headers: list[str]
    paths: list[str]
    preamble: list[list[Cell]]
    rows: list[tuple[int, list[Cell]]]

    def columns(self, aliases):
        return [i for i, h in enumerate(self.headers) if norm(h) in {norm(a) for a in aliases}]

    def col(self, *aliases):
        matches = self.columns(aliases)
        return matches[0] if matches else None

    def get(self, row, *aliases):
        c = self.col(*aliases)
        return row[c].value if c is not None else None


def read_tables(path, data_only=True):
    """One streaming pass per workbook, merging only header labels, never data."""
    try:
        wb = load_workbook(path, read_only=True, data_only=data_only)
        tables = []
        for ws in wb:
            iterator = ws.iter_rows()
            head = []
            for _ in range(min(ws.max_row or 0, 30)):
                row = next(iterator, None)
                if row is None:
                    break
                head.append([Cell(c.value, c.number_format or "") for c in row])
            scores = []
            keys = {
                "артикул",
                "код",
                "вид номенклатуры",
                "вид ассортимента",
                "наименование",
                "номенклатура",
                "начальный остаток",
            }
            for i, row in enumerate(head):
                score = sum(norm(c.value) in keys for c in row)
                scores.append((score, i))
            if not scores or max(scores)[0] < 2:
                continue
            _, hi = max(scores, key=lambda x: (x[0], -x[1]))
            headers = [str(c.value or "") for c in head[hi]]
            paths = list(headers)
            # Hierarchical month blocks: horizontal carry only on group-title rows.
            for upper in head[max(0, hi - 2) : hi]:
                carried = ""
                for c, value in enumerate(upper):
                    if value.value is not None:
                        carried = str(value.value)
                    if carried and any(word in norm(carried) for word in ("план по", "план markup")):
                        paths[c] = carried + " / " + headers[c]
                    elif value.value is not None:
                        paths[c] = str(value.value) + " / " + paths[c]
            rows = [(i + 1, r) for i, r in enumerate(head) if i > hi]
            rows.extend(
                (i, [Cell(c.value, c.number_format or "") for c in row])
                for i, row in enumerate(iterator, len(head) + 1)
            )
            tables.append(Table(ws.title, hi + 1, headers, paths, head[: hi + 1], rows))
        wb.close()
        if not tables:
            raise InputError("Не найдена таблица с заголовками артикула или вида номенклатуры.")
        return tables
    except InputError:
        raise
    except Exception as exc:
        raise InputError("Не удалось прочитать XLSX. Проверьте структуру и защиту файла.") from exc
