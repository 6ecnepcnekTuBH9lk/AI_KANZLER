"""Real inputs -> local HTTP -> stored article JSON -> XLSX; no external service."""

import hashlib
import json
import time
from contextlib import ExitStack
from pathlib import Path

import httpx
from merchandise_fixture import DIRECTORY, KEYS, NAMES, ROOT
from openpyxl import load_workbook

source = Path.home() / "Downloads"
destination = ROOT / "test-results/real-merchandise"
destination.mkdir(parents=True, exist_ok=True)
before = {name: hashlib.sha256((source / name).read_bytes()).hexdigest() for name in NAMES}
with httpx.Client(base_url="http://127.0.0.1:8001", timeout=180) as client:
    with ExitStack() as stack:
        files = [
            (
                "files",
                (
                    name,
                    stack.enter_context((source / name).open("rb")),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
            for name in NAMES
        ]
        response = client.post("/api/merchandise/runs", files=files)
        response.raise_for_status()
        identifier = response.json()["id"]
    print("Job:", identifier, flush=True)
    stop = time.monotonic() + 180
    while time.monotonic() < stop:
        response = client.get(f"/api/merchandise/runs/{identifier}")
        response.raise_for_status()
        job = response.json()
        if job["status"] in ("failed", "completed"):
            break
        time.sleep(0.5)
    assert job["status"] == "completed", job
    result = client.get(f"/api/merchandise/runs/{identifier}/summary").json()
    sections = {}
    for section in (
        "articles",
        "categories",
        "weekly-plan",
        "pricing",
        "second-wave",
        "actions",
        "data-quality",
    ):
        rows, offset = [], 0
        while True:
            response = client.get(
                f"/api/merchandise/runs/{identifier}/{section}", params={"limit": 5000, "offset": offset}
            )
            response.raise_for_status()
            page = response.json()
            rows += page["items"]
            if len(rows) >= page["total"]:
                break
            offset = len(rows)
        sections[section] = rows
    expected = json.loads((DIRECTORY / "expected.json").read_text(encoding="utf-8"))
    # API persistence must preserve every golden field, including nested second-wave data.
    assert [{k: a.get(k) for k in KEYS} for a in sections["articles"]] == expected
    response = client.get(f"/api/merchandise/runs/{identifier}/export")
    response.raise_for_status()
    path = destination / job["output_name"]
    path.write_bytes(response.content)
    wb = load_workbook(path, data_only=False)
    assert len(wb.sheetnames) == 8 and wb["Артикулы"].max_row == 339
    assert all(c.data_type != "f" for ws in wb for row in ws for c in row)
    wb.close()
after = {name: hashlib.sha256((source / name).read_bytes()).hexdigest() for name in NAMES}
assert before == after
(destination / "result.json").write_text(
    json.dumps({"result": result, "sections": sections}, ensure_ascii=False, indent=2), encoding="utf-8"
)
(destination / "http-verification.json").write_text(
    json.dumps(
        {
            "job_id": identifier,
            "inputs_sha256": before,
            "article_count": 338,
            "golden_matches": True,
            "output": str(path),
        },
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)
print(json.dumps(result["display_summary"], ensure_ascii=False, indent=2))
