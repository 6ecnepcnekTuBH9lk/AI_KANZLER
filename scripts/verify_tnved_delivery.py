"""Verify live output, original OOXML integrity and all 57 reviewed decisions; publish diffs."""

import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "backend/tests/fixtures/tnved"
OUT = ROOT / "test-results/real-tnved"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def main():
    baseline = json.loads((FIXTURE / "baseline.json").read_text(encoding="utf-8"))
    expected = json.loads((FIXTURE / "expected.json").read_text(encoding="utf-8"))
    live = json.loads((OUT / "result.json").read_text(encoding="utf-8"))
    items = live["sections"]["items"]
    assert len(items) == 57
    source = Path.home() / "Downloads" / baseline["name"]
    destination = OUT / (source.stem + "_с_кодами_ТНВЭД.xlsx")
    assert source.resolve() != destination.resolve()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == baseline["sha256"]
    before = baseline["app_before"]["sections"]["items"]
    historical_skill = baseline["previous_files"][1]["rows"]
    skill_path = source.parent / expected["skill_file"]
    if not skill_path.exists():
        skill_path = source.parent / "Коды ТНВЭД скилл.xlsx"
    skill_sha = hashlib.sha256(skill_path.read_bytes()).hexdigest()
    book = load_workbook(skill_path, data_only=False)
    values = list(book.active.values)
    skill = [dict(zip(values[0], row, strict=True)) for row in values[1:]]
    book.close()
    assert len(skill) == len(historical_skill) == 57
    # A refreshed user-format copy may change comments/source headers, never goods or codes.
    for current, historical in zip(skill, historical_skill, strict=True):
        for key in baseline["rows"][0]:
            assert str(current.get(key) or "") == str(historical.get(key) or ""), key
    followup_path = OUT / "followup-before.json"
    followup_before = (
        {r["row"]: r for r in json.loads(followup_path.read_text(encoding="utf-8"))}
        if followup_path.exists()
        else {}
    )
    rows, changed, differences = [], [], []
    for raw, actual, old, previous, target in zip(
        baseline["rows"], items, before, skill, expected["rows"], strict=True
    ):
        assert actual["row"] == target["row"]
        assert actual["code"] == target["code"], f"Breaking code regression at row {actual['row']}"
        if actual["code"]:
            assert re.fullmatch(r"[0-9]{10}", actual["code"])
            assert actual["evidence"]["code"] == actual["code"]
            assert actual["evidence"]["url"] == "https://www.alta.ru/tnved/code/" + actual["code"] + "/"
            assert actual["source_evidence"]
            assert sum(c["verdict"] == "match" for c in actual["candidates"]) == 1
            assert all(
                c["verdict"] == "excluded" for c in actual["candidates"] if c["code"] != actual["code"]
            )
        else:
            assert actual["status"] == "Требуется уточнение", actual["comment"]
            assert actual["reason_categories"] == ["MISSING_INPUT"]
            for fragment in target["required_reason_fragments"]:
                assert fragment in actual["comment"]
        skill_code = previous.get("Код ТНВЭД") or None
        skill_code = str(skill_code) if skill_code else None
        if skill_code != actual["code"]:
            assert target["skill_difference_category"] in ("A", "B", "C", "D", "E")
            differences.append(actual["row"])
        if old["code"] != actual["code"]:
            changed.append(actual["row"])
        rows.append(
            {
                "row": actual["row"],
                "source_code": actual["source_code"],
                "article": actual["article"],
                "nomenclature": raw["Номенклатура"],
                "name": raw["Наименование"],
                "features": actual["features"],
                "skill_result": skill_code,
                "app_before": old["code"],
                "app_after": actual["code"],
                "semantic_status": target["status"],
                "reason_category": target["category"],
                "missing_input": actual["missing_input"],
                "missing_rule": actual["missing_rule"],
                "classifier_limitation": actual["missing_rule"],
                "unresolved_reason": actual["comment"],
                "user_comment": actual["user_comment"],
                "source_url": actual["source_url"],
                "skill_comment": previous.get("Комментарий") or "",
                "skill_source": previous.get("Источник") or previous.get("Источник Alta.ru") or "",
                "followup_before": followup_before.get(actual["row"]),
                "alta_candidates": actual["candidates"],
                "evidence": actual["evidence"],
                "source_evidence": actual["source_evidence"],
                "difference_category": target["skill_difference_category"],
                "decision": target["reviewed_decision"],
            }
        )
    with ZipFile(source) as a, ZipFile(destination) as b:
        assert a.namelist() == b.namelist() and a.comment == b.comment
        changed_parts = [name for name in a.namelist() if a.read(name) != b.read(name)]
        assert set(changed_parts) <= {"xl/styles.xml", "xl/worksheets/sheet1.xml"}
        original, result = (etree.fromstring(z.read("xl/worksheets/sheet1.xml")) for z in (a, b))
        assert [r.get("r") for r in original.findall("s:sheetData/s:row", NS)] == [
            r.get("r") for r in result.findall("s:sheetData/s:row", NS)
        ]
        prior_cells = {c.get("r"): c for c in original.findall("s:sheetData/s:row/s:c", NS)}
        after_cells = {c.get("r"): c for c in result.findall("s:sheetData/s:row/s:c", NS)}
        # The export normalizes the requested code header as well as its values.
        permitted = {"F1"} | {"F" + str(i["row"]) for i in items}
        for coordinate, cell in prior_cells.items():
            if coordinate not in permitted or cell.find("s:f", NS) is not None:
                assert etree.tostring(cell) == etree.tostring(after_cells[coordinate]), coordinate
        for element in ("mergeCells", "sheetViews", "autoFilter", "sheetFormatPr", "drawing"):
            left, right = original.find("s:" + element, NS), result.find("s:" + element, NS)
            assert (etree.tostring(left) if left is not None else None) == (
                etree.tostring(right) if right is not None else None
            )
        # Original styles are an unchanged prefix; only result styles may be appended.
        prior_styles, after_styles = (etree.fromstring(z.read("xl/styles.xml")) for z in (a, b))
        old_xfs, new_xfs = prior_styles.find("s:cellXfs", NS), after_styles.find("s:cellXfs", NS)
        assert all(etree.tostring(x) == etree.tostring(new_xfs[i]) for i, x in enumerate(old_xfs))
    wb = load_workbook(destination, data_only=False)
    ws = wb.active
    assert ws.max_column == 13
    assert ws.cell(1, 12).value == "Комментарий"
    assert ws.cell(1, 13).value == "Источник"
    for actual in items:
        cell = ws.cell(actual["row"], 6)
        assert (cell.value or None) == actual["code"]
        if actual["code"]:
            assert cell.data_type == "s" and cell.number_format == "@"
        comment = actual["user_comment"]
        assert len(comment) <= 180 and comment.count(".") <= 2
        assert (ws.cell(actual["row"], 12).value or "") == comment
        url = f"https://www.alta.ru/tnved/code/{actual['code']}/" if actual["code"] else ""
        assert actual["source_url"] == url
        assert (ws.cell(actual["row"], 13).value or "") == url
        assert bool(comment) == (not actual["code"])
    wb.close()
    counts = Counter(r["reason_category"] for r in rows if r["semantic_status"] == "UNRESOLVED")
    summary = {
        "baseline_rows": 57,
        "original_empty_codes": 57,
        "skill_confirmed": 23,
        "skill_file_checked": str(skill_path),
        "skill_sha256_checked": skill_sha,
        "skill_historical_sha256": expected["skill_sha256"],
        "skill_goods_and_codes_match_historical": True,
        "app_before_confirmed": 13,
        "app_after_confirmed": sum(bool(r["app_after"]) for r in rows),
        "unresolved": sum(not r["app_after"] for r in rows),
        "reason_categories": dict(counts),
        "changed_from_app_before": changed,
        "different_from_skill": differences,
        "source_header": "Источник",
        "source_urls": sum(bool(r["source_url"]) for r in rows),
        "source_errors": sum(i["status"] == "Техническая ошибка" for i in items),
        "followup_changed_codes": [
            r["row"] for r in rows if r["followup_before"] and r["followup_before"]["code"] != r["app_after"]
        ],
        "skill_format_comparison": {
            "rows_checked": len(rows),
            "same_comment": sum(r["skill_comment"] == r["user_comment"] for r in rows),
            "same_source": sum(r["skill_source"] == r["source_url"] for r in rows),
            "comments_written": sum(bool(r["user_comment"]) for r in rows),
        },
        "shoe_confirmed": sum(bool(r["app_after"]) for r in rows if r["features"]["kind"] == "обувь"),
        "shoe_unresolved": sum(not r["app_after"] for r in rows if r["features"]["kind"] == "обувь"),
        "source_sha256": baseline["sha256"],
        "output_sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
        "changed_zip_parts": changed_parts,
        "ooxml_verified": True,
        "output_file": str(destination),
    }
    (OUT / "verification.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "differential-57.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    markdown = [
        "# ТН ВЭД: сравнение 57 строк",
        "",
        "Эталон Skill: файл с 23 кодами, подтверждён пользователем. Источник решений — актуальная Alta; сохранённые карточки используются только в offline regression.",
        f"Пользовательский формат сверяется с текущим файлом «{skill_path.name}». Все 57 товаров и коды совпали с историческим снимком; комментарии и оформление источника читаются из текущего файла. SHA256 текущего файла: {skill_sha}; исторического: {expected['skill_sha256']}.",
        "",
        "Точечная доработка: 17 → 19 подтверждений, 40 → 38 уточнений. Изменились только коды строк 18–19; остальные 17 подтверждений сохранены. Историческое «До» ниже относится к более раннему baseline из commit 787c689 (13 подтверждений).",
        "",
        "| Строка | Код / Артикул | Skill | До (исторический baseline) | После | Причина / решение |",
        "|---|---|---|---|---|---|",
    ]
    for row in rows:
        reason = ((row["difference_category"] + ": ") if row["difference_category"] else "") + row["decision"]
        markdown.append(
            f"| {row['row']} | {row['source_code']} / {row['article']} | {row['skill_result'] or '—'} | {row['app_before'] or '—'} | {row['app_after'] or '—'} | {reason} |"
        )
    markdown += [
        "",
        "Изменились относительно исторического baseline: " + ", ".join(map(str, changed)) + ".",
        "Отличаются от Skill: " + ", ".join(map(str, differences)) + ".",
        "",
        "A — отсутствовало правило; B — ошибка извлечения; C — не исследована ветвь; D — не хватает входных данных; E — прежнее недоказанное предположение. Четыре отличия от Skill относятся к E: назначение изделий в строках 24–27 не доказано.",
        "",
        "## Пользовательский формат: построчная сверка со Skill",
        "",
        "Во всех 57 строках сверены код, комментарий и источник. Комментарий запрашивает ближайший необходимый признак; полная проверка условий остаётся в UI/API. У UNRESOLVED источник намеренно пуст, даже если Skill приводил обзорную страницу группы. У CONFIRMED комментарий пуст, источник — прямая карточка кода.",
        "",
        "| Строка | Комментарий Skill | Комментарий Excel | Источник Skill | Источник Excel |",
        "|---|---|---|---|---|",
    ]

    def cell_text(value):
        return str(value or "—").replace("|", "\\|").replace("\n", " ")

    for row in rows:
        markdown.append(
            "| "
            + " | ".join(
                cell_text(row[k])
                for k in ("row", "skill_comment", "user_comment", "skill_source", "source_url")
            )
            + " |"
        )
    if followup_before:
        markdown += [
            "",
            "## Изменения относительно предыдущих 17 подтверждений",
            "",
            "Изменились коды в строках: " + ", ".join(map(str, summary["followup_changed_codes"])) + ".",
            "",
            "| Строка | Комментарий до | Комментарий после |",
            "|---|---|---|",
        ]
        for row in rows:
            markdown.append(
                "| "
                + " | ".join(
                    cell_text(v) for v in (row["row"], row["followup_before"]["comment"], row["user_comment"])
                )
                + " |"
            )
    (OUT / "skill-comparison-57.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
