from copy import copy, deepcopy
from types import SimpleNamespace
from zipfile import ZipFile

import httpx
import pytest
from app.core.exceptions import SourceUnavailable
from app.modules.tnved.alta import AltaClient
from app.modules.tnved.classifier import CandidateValidator, ClassificationEngine, valid_code
from app.modules.tnved.features import extract_features
from app.modules.tnved.service import run
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def shirt(**kwargs):
    return {"Номенклатура": "Рубашка", "Пол": "Мужской", "Трикотаж": "Да", "Состав": "100% хлопок", **kwargs}


def evidence(code="6105100000"):
    return {
        "code": code,
        "url": f"https://www.alta.ru/tnved/code/{code}/",
        "verified_at": "2026-09-07T00:00:00+00:00",
        "sha256": "abc",
        "description": "Рубашки мужские трикотажные из хлопчатобумажной пряжи",
        "levels": [
            {"code": "50-63", "description": "Текстиль"},
            {"code": "61", "description": "Одежда"},
            {
                "code": "6105",
                "description": "Рубашки трикотажные машинного или ручного вязания, мужские или для мальчиков",
            },
            {"code": code, "description": "из хлопчатобумажной пряжи"},
        ],
    }


def test_feature_signature_and_material():
    a = extract_features(shirt(Артикул="A", Цвет="Красный"))
    b = extract_features(shirt(Артикул="B", Цвет="Синий"))
    assert a.signature == b.signature and a.material == "хлопок" and a.knit is True
    c = extract_features(shirt(Состав="20% шерсть, 77% лиоцелл, 3% эластан; подкладка: 100% полиэстер"))
    assert c.material == "искусственные" and c.composition == [
        ("искусственные", 77.0),
        ("синтетические", 3.0),
        ("шерсть", 20.0),
    ]
    assert not c.missing


def test_material_reverse_order_and_ambiguity():
    assert extract_features(shirt(Состав="хлопок 100%")).material == "хлопок"
    # Section XI 2(A,B): equal chapter weights select the last chapter (54/55).
    assert extract_features(shirt(Состав="50% хлопок, 50% полиэстер")).material == "синтетические"
    assert extract_features(shirt(Состав="натуральная кожа")).missing


def test_knit_conflict_and_shoe_strictness():
    f = extract_features(shirt(Наименование="Рубашка трикотажная", Трикотаж="Нет"))
    assert any("противореч" in x for x in f.missing)
    f = extract_features({"Номенклатура": "Обувь", "Пол": "Мужской", "Состав": "натуральная кожа"})
    # Insole length is requested by a candidate, not unconditionally for all shoes.
    assert "материал подошвы" in f.missing and "материал верха" in f.missing


@pytest.mark.parametrize(
    "code",
    ["61", "6105", "610510", "61051000", "6105 100000", "6105100000, 6105201000", "рубашки", 6105100000],
)
def test_reject_incomplete_codes(code):
    assert not valid_code(code)


def test_candidate_verification_and_unresolved():
    validator = CandidateValidator()
    f = extract_features(shirt())
    assert validator.validate(f, evidence()) == "match"
    wrong = deepcopy(evidence())
    wrong["levels"][-1]["description"] = "из синтетических нитей"
    assert validator.validate(f, wrong) == "excluded"
    unclear = deepcopy(evidence())
    unclear["levels"].append({"code": "6105100000", "description": "прочие, массой более 600 г"})
    assert validator.validate(f, unclear) == "unresolved"
    wrong_gender = deepcopy(evidence())
    wrong_gender["levels"][2]["description"] = "Рубашки трикотажные для женщин"
    assert validator.validate(f, wrong_gender) == "excluded"


def test_alta_heading_with_trimmed_zero():
    from app.modules.tnved.alta import CandidateFinder

    class TreeClient:
        def get(self, path, params):
            if params == {"tnved": "6110"}:
                return "<ul></ul>"
            assert params == {"tnved": "611"}
            return '<ul><li data-source="611" data-uin="root"><ul><li data-source="6110209100" data-uin="leaf"><a href="/tnved/code/6110209100/">код</a></li></ul></li></ul>'

    finder = CandidateFinder(TreeClient())
    f = extract_features(shirt(Номенклатура="Джемпер"))
    assert finder.find(f) == ["6110209100"]


def test_multiple_candidates_no_guess_and_grouping():
    client = SimpleNamespace(verify=lambda c: evidence(c))
    engine = ClassificationEngine(client)
    engine.finder = SimpleNamespace(find=lambda f: ["6105100000", "6105100001"])
    result = engine.classify(extract_features(shirt()))
    assert result["code"] is None and result["comment"]
    engine.finder = SimpleNamespace(find=lambda f: pytest.fail("Same signature must use memo"))
    assert engine.classify(extract_features(shirt(Цвет="red"))) == result


def test_unavailable_alta_not_guessed():
    class Unavailable:
        def find(self, f):
            raise SourceUnavailable("Не удалось проверить код на Alta.ru: сайт недоступен.")

    engine = ClassificationEngine(None)
    engine.finder = Unavailable()
    result = engine.classify(extract_features(shirt()))
    assert result["code"] is None and result["status"] == "Техническая ошибка"


def test_http_blocks_and_html_changes():
    client = AltaClient(httpx.MockTransport(lambda r: httpx.Response(403)), interval=0)
    with pytest.raises(SourceUnavailable):
        client.verify("6105100000")
    client.close()
    client = AltaClient(
        httpx.MockTransport(lambda r: httpx.Response(200, text="<html>Changed structure</html>")), interval=0
    )
    with pytest.raises(SourceUnavailable):
        client.verify("6105100000")
    client.close()


def make_input(path, success_only=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Номенклатура"
    ws.append(
        ["Артикул", "Номенклатура", "Пол", "Трикотаж", "Состав", "Код ТН ВЭД", "Примечание", "Комментарий"]
    )
    ws.append(["A", "Рубашка", "Мужской", "Да", "100% хлопок", None, "=1+2", None])
    if not success_only:
        ws.append(
            ["B", "Обувь", "Мужской", "Нет", "натуральная кожа", None, "Сохранить", "Текст пользователя"]
        )
        ws.append(["C", "Рубашка", "Мужской", "Да", "100% хлопок", "1234", "Существующий код", None])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 24
    ws.row_dimensions[2].height = 31
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["F2"].fill = PatternFill("solid", fgColor="FFAA11")
    other = wb.create_sheet("Другой лист")
    other.merge_cells("A1:C1")
    other["A1"] = "Сохранить"
    other["D2"] = "=SUM(1,2)"
    wb.save(path)


def test_excel_preserves_all_parts_values_styles_existing_and_comments(tmp_path, monkeypatch):
    path = tmp_path / "input.xlsx"
    make_input(path)

    def classify(self, f):
        if f.kind == "обувь":
            return {
                "code": None,
                "status": "Требуется уточнение",
                "comment": "Не указан материал подошвы.",
                "evidence": None,
                "candidates": [],
            }
        return {
            "code": "6105100000",
            "status": "Код определён",
            "comment": "",
            "evidence": evidence(),
            "candidates": [],
        }

    monkeypatch.setattr(ClassificationEngine, "classify", classify)
    file = SimpleNamespace(path=path, original_name=path.name)
    result, _sections, out, *_ = run([file], {}, tmp_path / "out", lambda p, m: None, None)
    assert result["total"] == 3 and result["counts"]["Уже имел код"] == 1
    before = load_workbook(path)
    after = load_workbook(out)
    assert after.sheetnames == before.sheetnames
    a = after.active
    b = before.active
    assert a["F2"].value == "6105100000" and a["F2"].data_type == "s" and a["F2"].number_format == "@"
    assert a["F4"].value == "1234" and a["G2"].value == "=1+2"
    assert _sections["items"][1]["original_comment"] == "Текст пользователя"
    assert a["H3"].value == _sections["items"][1]["user_comment"] and "подошвы" in a["H3"].value
    assert copy(a["A2"].font) == copy(b["A2"].font) and copy(a["F2"].fill) == copy(b["F2"].fill)
    assert a.freeze_panes == b.freeze_panes and a.auto_filter.ref == b.auto_filter.ref
    assert a.column_dimensions["A"].width == 24 and a.row_dimensions[2].height == 31
    with ZipFile(path) as z1, ZipFile(out) as z2:
        assert z1.read("xl/worksheets/sheet2.xml") == z2.read("xl/worksheets/sheet2.xml")
    before.close()
    after.close()


def test_output_columns_exist_even_when_all_succeed(tmp_path, monkeypatch):
    path = tmp_path / "input.xlsx"
    make_input(path, True)
    wb = load_workbook(path)
    wb.active.delete_cols(8)
    wb.save(path)
    monkeypatch.setattr(
        ClassificationEngine,
        "classify",
        lambda self, f: {
            "code": "6105100000",
            "status": "Код определён",
            "comment": "",
            "evidence": evidence(),
            "candidates": [],
        },
    )
    _, _, out, *_ = run(
        [SimpleNamespace(path=path, original_name=path.name)], {}, tmp_path / "out", lambda p, m: None, None
    )
    wb = load_workbook(out)
    assert wb.active.max_column == 9
    assert wb.active["H1"].value == "Комментарий" and not wb.active["H2"].value
    assert wb.active["I1"].value == "Источник"
    assert wb.active["I2"].value == "https://www.alta.ru/tnved/code/6105100000/"
    wb.close()


def test_formula_in_code_and_comment_is_preserved(tmp_path, monkeypatch):
    path = tmp_path / "formulas.xlsx"
    make_input(path)
    wb = load_workbook(path)
    wb.active["F2"] = '=TEXT(6105100000,"0")'
    wb.active["H3"] = '="Комментарий пользователя"'
    wb.save(path)
    monkeypatch.setattr(
        ClassificationEngine,
        "classify",
        lambda self, f: {
            "code": None,
            "status": "Требуется уточнение",
            "comment": "Не указан материал подошвы.",
            "evidence": None,
            "candidates": [],
        },
    )
    _, _, out, *_ = run(
        [SimpleNamespace(path=path, original_name=path.name)], {}, tmp_path / "out", lambda p, m: None, None
    )
    wb = load_workbook(out)
    assert wb.active["F2"].value == '=TEXT(6105100000,"0")'
    assert wb.active["H3"].value == '="Комментарий пользователя"'
    assert wb.active["I1"].value == "Комментарий ТН ВЭД"
    assert "подошвы" in wb.active["I3"].value
    wb.close()
    # Reprocessing must reuse the dedicated comment column beside the formula,
    # not create a new identical header on every run.
    _, _, again, *_ = run(
        [SimpleNamespace(path=out, original_name=out.name)], {}, tmp_path / "again", lambda p, m: None, None
    )
    wb = load_workbook(again)
    assert wb.active.max_column == 10
    assert wb.active["I3"].value.count("материал подошвы") == 1
    assert wb.active["H3"].value == '="Комментарий пользователя"'
    wb.close()


def test_explicit_verification_preserves_formula_and_reports_unproven(tmp_path, monkeypatch):
    path = tmp_path / "formula_verify.xlsx"
    make_input(path, True)
    wb = load_workbook(path)
    wb.active["F2"] = '=TEXT(6105100000,"0")'
    wb.save(path)
    wb.close()
    monkeypatch.setattr(
        ClassificationEngine,
        "classify",
        lambda self, f: {
            "code": "6105100000",
            "status": "Код определён",
            "comment": "",
            "evidence": evidence(),
            "candidates": [],
        },
    )
    _, sections, out, *_ = run(
        [SimpleNamespace(path=path, original_name=path.name)],
        {"verify_existing": True},
        tmp_path / "out",
        lambda p, m: None,
        None,
    )
    item = sections["items"][0]
    assert item["code"] is None and item["proposed_code"] == "6105100000"
    assert item["status"] == "Исходный код не подтверждён"
    assert item["existing_verification"]["status"] == "UNPROVEN"
    wb = load_workbook(out)
    assert wb.active["F2"].value == '=TEXT(6105100000,"0")'
    assert "формула сохранена" in wb.active["H2"].value
    wb.close()


def test_existing_code_verified_separately_not_erased_when_unproven(tmp_path, monkeypatch):
    path = tmp_path / "existing_verify.xlsx"
    make_input(path, True)
    wb = load_workbook(path)
    wb.active["F2"] = "6105100000"
    wb.save(path)
    wb.close()
    checked = []
    client = SimpleNamespace(verify=lambda code: checked.append(code) or evidence(code), close=lambda: None)
    monkeypatch.setattr(
        ClassificationEngine,
        "classify",
        lambda self, f: {
            "code": None,
            "status": "Требуется уточнение",
            "comment": "Не исключён второй кандидат.",
            "evidence": None,
            "candidates": [],
        },
    )
    _, sections, out, *_ = run(
        [SimpleNamespace(path=path, original_name=path.name)],
        {"verify_existing": True},
        tmp_path / "out",
        lambda p, m: None,
        None,
        client_factory=lambda: client,
    )
    assert checked == ["6105100000"]
    assert sections["items"][0]["existing_verification"]["status"] == "UNPROVEN"
    wb = load_workbook(out)
    assert wb.active["F2"].value == "6105100000"
    assert "не подтверждён" in wb.active["H2"].value
    wb.close()


def test_xlsx_drawing_and_custom_parts_survive(tmp_path, monkeypatch):
    import base64

    from lxml import etree

    path = tmp_path / "drawings.xlsx"
    make_input(path)
    with ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    sheet = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
    drawing = etree.SubElement(sheet, "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}drawing")
    drawing.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rIdImage")
    parts["xl/worksheets/sheet1.xml"] = etree.tostring(sheet)
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>'
    )
    parts["xl/drawings/drawing1.xml"] = (
        b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><xdr:oneCellAnchor><xdr:from><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="76200" cy="76200"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="1" name="Image 1"/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr><a:prstGeom prst="rect"/></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>'
    )
    parts["xl/drawings/_rels/drawing1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>'
    )
    parts["xl/media/image1.png"] = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aF1kAAAAASUVORK5CYII="
    )
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
        Extension="png",
        ContentType="image/png",
    )
    etree.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        PartName="/xl/drawings/drawing1.xml",
        ContentType="application/vnd.openxmlformats-officedocument.drawing+xml",
    )
    parts["[Content_Types].xml"] = etree.tostring(content_types)
    parts["customXml/item1.xml"] = b"<custom>keep</custom>"
    with ZipFile(path, "w") as z:
        for name, content in parts.items():
            z.writestr(name, content)
        z.comment = b"original archive comment"
    monkeypatch.setattr(
        ClassificationEngine,
        "classify",
        lambda self, f: {
            "code": None,
            "status": "Требуется уточнение",
            "comment": "Не указан материал подошвы.",
            "evidence": None,
            "candidates": [],
        },
    )
    _, _, out, *_ = run(
        [SimpleNamespace(path=path, original_name=path.name)], {}, tmp_path / "out", lambda p, m: None, None
    )
    with ZipFile(path) as before, ZipFile(out) as after:
        assert before.namelist() == after.namelist() and before.comment == after.comment
        for name in before.namelist():
            if name not in ("xl/styles.xml", "xl/worksheets/sheet1.xml"):
                assert before.read(name) == after.read(name), name
    wb = load_workbook(out)
    assert wb.active["G2"].value == "=1+2"
    wb.close()


def test_empty_name_falls_back_to_nomenclature(tmp_path, monkeypatch):
    path = tmp_path / "empty_name.xlsx"
    wb = Workbook()
    wb.active.append(["Наименование", "Номенклатура", "Пол", "Состав", "Трикотаж"])
    wb.active.append([None, "Рубашка", "Мужской", "100% хлопок", "Да"])
    wb.save(path)
    wb.close()
    monkeypatch.setattr(
        ClassificationEngine,
        "classify",
        lambda self, f: {
            "code": None,
            "status": "Требуется уточнение",
            "comment": "Источник не проверен.",
            "evidence": None,
            "candidates": [],
        },
    )
    result, _, _, *_ = run(
        [SimpleNamespace(path=path, original_name=path.name)], {}, tmp_path / "out", lambda p, m: None, None
    )
    assert result["total"] == 1
