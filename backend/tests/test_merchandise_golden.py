"""Permanent article-level regression plus a separate arithmetic oracle.

The oracle does not call planning/metrics/decisions. Parsing is covered separately.
Golden recommendations are reviewed snapshots, not an independent expert verdict.
"""

import json
from calendar import monthrange
from datetime import date, datetime, timedelta
from statistics import mean

import pytest
from app.modules.merchandise import importers, presentation, reporting
from app.modules.merchandise.service import run
from app.shared.excel import norm
from openpyxl import load_workbook

from scripts.merchandise_fixture import DIRECTORY, materialize, snapshot


@pytest.fixture(scope="module")
def golden_run(tmp_path_factory):
    directory = tmp_path_factory.mktemp("golden")
    files = materialize(directory / "inputs")
    result = run(files, {}, directory / "output", lambda *_: None)
    return files, result


def compare(actual, expected):
    if isinstance(expected, dict):
        assert actual.keys() == expected.keys()
        for key in expected:
            compare(actual[key], expected[key])
    elif isinstance(expected, list):
        assert len(actual) == len(expected)
        for a, e in zip(actual, expected):
            compare(a, e)
    elif isinstance(expected, float):
        assert actual == pytest.approx(expected, rel=1e-10, abs=1e-9)
    else:
        assert actual == expected


def test_every_real_article_matches_golden(golden_run):
    _, (_, sections, *_) = golden_run
    expected = json.loads((DIRECTORY / "expected.json").read_text(encoding="utf-8"))
    actual = snapshot(sections["articles"])
    assert len(actual) == len(expected) == 338
    for a, e in zip(actual, expected):
        try:
            compare(a, e)
        except AssertionError as error:
            raise AssertionError(f"Golden mismatch: {a['article']}") from error


def test_previous_arithmetic_is_unchanged(golden_run):
    _, (_, sections, *_) = golden_run
    previous = json.loads((DIRECTORY / "arithmetic_baseline.json").read_text(encoding="utf-8"))
    for a, before in zip(sections["articles"], previous):
        compare({key: a[key] for key in before}, before)


def test_real_articles_independent_arithmetic(golden_run):
    files, (result, sections, *_) = golden_run
    roles = importers.detect_roles(files)
    plans = {p["article"]: p for p in importers.read_plans(roles["article"][1])}
    cats = {
        (norm(p["kind"]), norm(p["assortment"])): p for p in importers.read_plans(roles["category"][1], True)
    }
    facts = {f["article"]: f for f in importers.read_fact(roles["fact"][1], "ОЗ26", set(plans))}
    control = date.fromisoformat(result["control_date"])
    start, end = date(2026, 9, 1), date(2027, 2, 28)
    computed = 0
    for a in sections["articles"]:
        f = facts[a["article"]]
        p = plans.get(a["article"])
        c = cats.get((norm(f["kind"]), norm(f["assortment"])))
        base = f["base"]
        assert a["base"] == base
        if base and base > 0:
            target = p["target"] if p and p["target"] is not None else c["target"] if c else None
            shape = p if p and any(v is not None for v in p["months"].values()) else c
            if target and shape and all(shape["months"].get(m) is not None for m in (9, 10, 11, 12, 1, 2)):
                denominator = sum(shape["months"][m] for m in (9, 10, 11, 12, 1, 2))
                pct = sum(
                    target
                    * shape["months"][(start + timedelta(days=i)).month]
                    / denominator
                    / monthrange((start + timedelta(days=i)).year, (start + timedelta(days=i)).month)[1]
                    for i in range((min(control, end) - start).days + 1)
                )
                assert a["target"] == pytest.approx(target)
                assert a["season_plan"] == pytest.approx(base * target)
                assert a["plan_pct"] == pytest.approx(pct)
                assert a["plan_units"] == pytest.approx(base * pct)
                computed += 1
        # The fixture control date is Sunday of W36: six official days of the boundary week.
        assert control == date(2026, 9, 6)
        boundary = f["weekly"].get(date(2026, 8, 31))
        official = boundary * 6 / 7 if boundary is not None else None
        assert (
            a["season_fact"] == pytest.approx(official) if official is not None else a["season_fact"] is None
        )
        prior = [v for d, v in f["weekly"].items() if d < date(2026, 8, 31)]
        known = [v for v in prior if v is not None] + ([boundary / 7] if boundary is not None else [])
        assert a["preseason"] == pytest.approx(sum(known)) if known else a["preseason"] is None
        if official is not None and base and base > 0:
            assert a["st"] == pytest.approx(official / base)
            if a["plan_units"]:
                assert a["execution"] == pytest.approx(official / a["plan_units"])
                assert a["gap_pp"] == pytest.approx(100 * (official / base - a["plan_pct"]))
                assert a["gap_units"] == pytest.approx(official - a["plan_units"])
        recent = []
        for i in range(4):
            value = f["weekly"].get(date(2026, 8, 31) - timedelta(days=7 * i))
            if value is None:
                break
            recent.append(value)
        pace = mean(recent) if len(recent) >= 2 else None
        assert a["pace"] == pytest.approx(pace) if pace is not None else a["pace"] is None
        if pace and f["stock"] is not None:
            assert a["cover"] == pytest.approx(f["stock"] / pace)
        if a["deadline"] and official is not None and pace is not None:
            left = max(0, (date.fromisoformat(a["deadline"]) - control).days / 7)
            assert a["forecast"] == pytest.approx(official + pace * left)
            if base:
                assert a["forecast_st"] == pytest.approx(a["forecast"] / base)
            if a["season_plan"] is not None and left > 0:
                assert a["required"] == pytest.approx(max(0, a["season_plan"] - official) / left)
        if wave := a["second_wave"]:
            if base and base > 0:
                assert wave["target70"] == pytest.approx(base * 0.7)
            if wave["first_sales"] is not None and pace is not None and wave["weeks"] is not None:
                assert wave["forecast"] == pytest.approx(wave["first_sales"] + pace * wave["weeks"])
                if wave["target70"] is not None:
                    assert wave["gap"] == pytest.approx(wave["forecast"] - wave["target70"])
    assert computed == 223


def test_full_real_fixture_first_batch_invariance(golden_run, tmp_path):
    files, (_, original, *_) = golden_run
    # Change only the copied ARTICLE input, never the permanent fixture or user's files.
    article_file = importers.detect_roles(files)["article"][0]
    wb = load_workbook(article_file.path)
    ws = wb.active
    columns = [c.column for c in ws[3] if c.value == "1-я партия"]
    assert columns
    for ri in range(4, ws.max_row + 1):
        for ci in columns:
            ws.cell(ri, ci, 10**12)
    path = tmp_path / article_file.original_name
    wb.save(path)
    wb.close()
    from types import SimpleNamespace

    copied = [
        SimpleNamespace(path=path, original_name=f.original_name) if f is article_file else f for f in files
    ]
    _, changed, *_ = run(copied, {}, tmp_path / "output", lambda *_: None)
    compare(snapshot(changed["articles"]), snapshot(original["articles"]))


def test_all_exported_values_and_compact_layout(golden_run):
    _, (result, sections, path, *_) = golden_run
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == list(reporting.SHEETS.values())
    for section, title in reporting.SHEETS.items():
        ws = wb[title]
        if section == "summary":
            records, columns = result["summary_rows"], reporting.SUMMARY_COLUMNS
        elif section == "data-quality":
            records, columns = presentation.group_quality(sections[section]), reporting.QUALITY_GROUP_COLUMNS
        elif section == "weekly-plan":
            records = sections[section]
            columns = [
                ("article", "Артикул"),
                ("code", "Код"),
                ("base", "Начальный остаток"),
                ("target", "Целевой процент реализации сезона, %"),
            ]
            columns += [(p["week_start"], p["label"]) for p in result["week_columns"]]
        else:
            records, columns = sections[section], reporting.SCHEMAS[section]
        assert ws.max_row == len(records) + 1
        assert [c.value for c in ws[1]] == [label for _, label in columns]
        assert ws.freeze_panes and ws.auto_filter.ref
        assert all(d.width <= (80 if section == "methodology" else 45) for d in ws.column_dimensions.values())
        assert all(
            d.height <= (48 if section == "methodology" else 42)
            for d in ws.row_dimensions.values()
            if d.height
        )
        for cells, record in zip(ws.iter_rows(min_row=2), records):
            for cell, (key, _) in zip(cells, columns):
                assert cell.data_type != "f" and cell.font.sz == 9
                expected = record.get(key)
                if key in presentation.UNIT_KEYS or (
                    section == "summary" and key == "value" and record.get("format") == "units"
                ):
                    expected = presentation.display_units(expected)
                    assert cell.number_format == "#,##0"
                if expected == "":
                    expected = None
                actual = cell.value
                if isinstance(actual, datetime):
                    actual = actual.date().isoformat()
                compare(actual, expected)
    assert not any("План неделя" in c.value for c in wb["Артикулы"][1])
    assert 32 <= wb["Артикулы"].max_column <= 38
    assert wb["Действия"].max_column == 8 and wb["Сводка"].max_column == 4
    wb.close()
