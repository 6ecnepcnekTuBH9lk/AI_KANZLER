"""Read-only resource inventory used during requirements analysis."""

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
source = root / "backend/app/resources/kanzler-merchandise-performance-controller"
out = root / "docs/source-review"
out.mkdir(parents=True, exist_ok=True)
doc = Document(source / "Управление ассортиментом.docx")
text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    text += "\n\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
(out / "corporate-rules.txt").write_text(text, encoding="utf-8")
inventory = []
for path in source.glob("*.xlsx"):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb:
        rows = []
        nonempty = 0
        for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = {str(i + 1): str(v) for i, v in enumerate(row) if v is not None}
            if cells:
                nonempty += 1
                if len(rows) < 18:
                    rows.append({"row": ri, "cells": cells})
        sheets.append(
            {
                "name": ws.title,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "nonempty_rows": nonempty,
                "sample": rows,
            }
        )
    inventory.append({"file": path.name, "sheets": sheets})
    wb.close()
(out / "workbooks.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
print(text)
print("\nWORKBOOK SUMMARY")
for book in inventory:
    print(book["file"], [(s["name"], s["rows"], s["columns"]) for s in book["sheets"]])
